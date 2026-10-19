"""Tests du screener : moteur point-in-time, filtres, portefeuille, métriques, puis test complet simulé."""
import datetime as dt
import os
import shutil
import sys
import tempfile
import time

import numpy as np
import pandas as pd

ICI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(ICI))
sys.path.insert(0, ICI)

import screener_sp500 as sc          # noqa: E402
import donnees_simulees as ds        # noqa: E402

SORTIE = os.path.join(tempfile.mkdtemp(prefix="test_moteur_"), "sortie")
OK, KO = [], []


def verifier(nom, cond, detail=""):
    (OK if cond else KO).append(nom)
    print(("  OK  " if cond else "  ÉCHEC ") + nom + (f"  ({detail})" if detail else ""))


D = ds.construire()

# ---------------------------------------------------------------------------
print("\n1. Moteur point-in-time (TTM, trimestres, première publication)")
cik = 1000
faits = sc.extraire_faits(D["jsons"][cik])
table = sc.table_fondamentale(faits)
ver = D["verites"][cik]
tri, q = ver["tri"], ver["q"]
fins = [pd.Timestamp(r[3]) for r in tri]
verifier("table non vide", len(table) > 40, f"{len(table)} lignes")

erreurs = []
for _, ligne in table.iterrows():
    k = fins.index(ligne["fin"])
    attendu = q["rn"][k - 3:k + 1].sum()
    if np.isfinite(ligne["resultat_net"]):
        erreurs.append(abs(ligne["resultat_net"] / attendu - 1))
verifier("résultat net TTM = somme des 4 trimestres", len(erreurs) > 30 and max(erreurs) < 1e-9,
         f"{len(erreurs)} lignes, écart max {max(erreurs):.2e}")

erreurs_cf = []
for _, ligne in table.iterrows():
    k = fins.index(ligne["fin"])
    attendu = q["cfo"][k - 3:k + 1].sum()
    if np.isfinite(ligne["cfo"]):
        erreurs_cf.append(abs(ligne["cfo"] / attendu - 1))
verifier("flux de trésorerie TTM (cumuls uniquement) = somme des 4 trimestres",
         len(erreurs_cf) > 30 and max(erreurs_cf) < 1e-9, f"écart max {max(erreurs_cf):.2e}")

rev_err = []
for _, ligne in table.iterrows():
    k = fins.index(ligne["fin"])
    if np.isfinite(ligne["revenus"]):
        rev_err.append(abs(ligne["revenus"] / q["rev"][k - 3:k + 1].sum() - 1))
verifier("chiffre d'affaires TTM malgré le changement d'étiquette en 2018",
         len(rev_err) > 30 and max(rev_err) < 1e-9, f"écart max {max(rev_err):.2e}")

# trimestre seul (T4 = exercice - 9 mois) via la surprise : vérifier q_rn indirectement
t4 = table[table["annuel"]]
verifier("lignes annuelles détectées (10-K)", len(t4) >= 15, f"{len(t4)}")
k = fins.index(t4["fin"].iloc[5])
attendu_surprise = (q["rn"][k] - q["rn"][k - 4]) / ver["bilan"]["Assets"][k]
verifier("surprise T4 = (T4 - T4 an passé) / actif", abs(table.loc[t4.index[5], "surprise_rn"] - attendu_surprise) < 1e-12)

# Première publication : la société 1002 a retraité son résultat 2014 dans le 10-K 2015
f2 = sc.extraire_faits(D["jsons"][1002])
t2 = sc.table_fondamentale(f2)
tri2, q2 = D["verites"][1002]["tri"], D["verites"][1002]["q"]
fins2 = [pd.Timestamp(r[3]) for r in tri2]
ligne14 = t2[(t2["fin"].dt.year == 2014) & t2["annuel"]].iloc[0]
k14 = fins2.index(ligne14["fin"])
verifier("valeur retraitée ignorée (première publication conservée)",
         abs(ligne14["resultat_net"] / q2["rn"][k14 - 3:k14 + 1].sum() - 1) < 1e-9)

# Exercice décalé + semaines de 52/53 (société 1001)
t1 = sc.table_fondamentale(sc.extraire_faits(D["jsons"][1001]))
verifier("exercice clos en septembre, trimestres de 52/53 semaines : TTM calculés",
         t1["resultat_net"].notna().mean() > 0.9, f"{t1['resultat_net'].notna().mean():.0%}")

# Banque : pas de bilan courant ni de marge brute
tb = sc.table_fondamentale(sc.extraire_faits(D["jsons"][1030]))
recent = tb[tb["fin"] >= "2011-06-30"]
verifier("banque : F-score calculé sur 7 signaux (une fois l'historique disponible)",
         set(recent["fscore_n"].dropna().unique()) == {7.0}, f"{sorted(tb['fscore_n'].dropna().unique())}")
verifier("F-score entre 0 et 9", table["fscore"].dropna().between(0, 9).all())

# Point-in-time : avant le dépôt du 10-K, on doit voir le trimestre précédent
pan = sc.panneau_fondamental({cik: table})
ligne_k = table[table["annuel"]].iloc[8]
avant = ligne_k["dispo"] - pd.Timedelta(days=1)
apres = ligne_k["dispo"] + pd.Timedelta(days=2)
x = sc.fondamentaux_aux_dates(pan, [avant, apres], [cik])
verifier("pas d'anticipation : la veille du dépôt, le 10-K n'est pas visible",
         x.iloc[0]["fin"] < ligne_k["fin"] and x.iloc[1]["fin"] == ligne_k["fin"])

# Nouvelle holding (nouveau CIK) : l'historique de l'ancienne société est repris, à l'identique
faits3 = sc.extraire_faits(D["jsons"][1003])
coupure = pd.Timestamp("2018-01-01")
parts = {9001: faits3[faits3["depot"] < coupure], 9002: faits3[faits3["depot"] >= coupure]}


class SECHolding:
    def faits(self, c):
        return parts[int(c)].copy()


coupure2 = pd.Timestamp("2022-01-01")       # deux réorganisations successives (Avago -> Broadcom Ltd -> Inc.)
parts = {9001: faits3[faits3["depot"] < coupure], 9002: faits3[(faits3["depot"] >= coupure) & (faits3["depot"] < coupure2)],
         9003: faits3[faits3["depot"] >= coupure2]}
sc.PREDECESSEURS.update({9003: 9002, 9002: 9001})
seul = sc.table_fondamentale(parts[9003])
fusion = sc.table_fondamentale(sc.faits_societe(SECHolding(), 9003))
ref = sc.table_fondamentale(faits3)
verifier("nouvelle holding : historique des sociétés précédentes repris à l'identique (2 réorganisations)",
         len(seul) < len(ref) and len(fusion) == len(ref)
         and np.allclose(fusion["fscore"].values, ref["fscore"].values, equal_nan=True)
         and np.allclose(fusion["revenus"].values, ref["revenus"].values, equal_nan=True),
         f"{len(seul)} lignes seule, {len(fusion)} avec l'historique, {len(ref)} attendues")
del sc.PREDECESSEURS[9003], sc.PREDECESSEURS[9002]

# ---------------------------------------------------------------------------
print("\n2. Filtres, scores, signaux, portefeuille")
cfg = sc.Config(email="test@domaine-test.org", dossier=SORTIE, n_positions=6, seuil_achat=80)
idx = [f"X{i}" for i in range(12)]
df = pd.DataFrame(index=pd.Index(idx, name="ticker"))
df["secteur"] = ["Industrials"] * 6 + ["Financials"] * 6
df["nom"] = idx
df["cik"] = np.arange(12, dtype=float)
df["historique_ok"] = True
df["vol"] = np.linspace(0.15, 0.40, 12)
df["mom"] = np.linspace(-0.2, 0.4, 12)
df["beta"] = 1.0
df["mdd"] = -0.2
for c in sc.COLONNES_FOND:
    df[c] = np.nan
df["fscore"] = [8, 3, 7, 7, 7, 7, 6, 6, 6, 6, 6, 6]
df["fscore_n"] = [9, 9, 9, 9, 9, 9, 7, 7, 7, 7, 7, 7]
df["actif"] = 1e10
df["couverture"] = [10, 10, 1.0, 10, 10, 10, 0.5, 0.5, 0.5, 0.5, 0.5, 0.5]
df["ebitda"] = 1e9
df["dette_nette"] = 1e9
df["dette_ebitda"] = 1.0
df["fcf_3a"] = [1e9, 1e9, 1e9, -1e8, 1e9, 1e9] + [-1e9] * 6
df["fcf"] = 1e9
df["accruals"] = 0.0
df["beneish"] = -3.0
f = sc.appliquer_filtres(df, cfg)
verifier("F-score 3/9 exclu", bool(f.at["X1", "echec_filtre"]) and "F-score" in f.at["X1", "motif"])
verifier("couverture 1x exclue (non financière)", bool(f.at["X2", "echec_filtre"]))
verifier("FCF négatif sur 3 ans exclu", bool(f.at["X3", "echec_filtre"]))
verifier("banques : couverture et FCF ignorés", not f.loc["X6":"X9", "echec_filtre"].any()
         and "Volatilité" in f.at["X10", "motif"] and "Couverture" not in f.at["X10", "motif"])
verifier("10 % les plus volatils exclus", bool(f.at["X11", "echec_filtre"]) and "Volatilité" in f.at["X11", "motif"])
verifier("F-score des banques ramené sur 9", abs(f.at["X6", "fscore_9"] - 6 * 9 / 7) < 1e-9)
dfn = df.copy()
dfn.loc[["X0", "X6"], "fscore_n"] = 5
dfn.loc[["X0", "X6"], "fscore"] = 4
fn5 = sc.appliquer_filtres(dfn, cfg)
verifier("F-score sur 5 critères : suffisant pour une financière (5 sur 7), pas pour une autre (6 sur 9)",
         abs(fn5.at["X6", "fscore_9"] - 4 * 9 / 5) < 1e-9 and np.isnan(fn5.at["X0", "fscore_9"])
         and not fn5.at["X0", "donnees_ok"] and fn5.at["X6", "donnees_ok"])
dfa = df.copy()
dfa["accruals"] = [-0.05, -0.04, -0.03, -0.02, -0.01, -0.005] + [0.0] * 6
fa = sc.appliquer_filtres(dfa, cfg)
dfa.loc["X5", "accruals"] = 0.06
fb = sc.appliquer_filtres(dfa, cfg)
verifier("accruals négatifs jamais éliminatoires, même parmi les pires du secteur",
         fa.at["X5", "f_accruals"] == "ok" and "Accruals" not in fa.at["X5", "motif"]
         and fa.at["X5", "s_accruals"] == 0.0, fa.at["X5", "motif"])
verifier("accruals positifs parmi les pires du secteur exclus",
         fb.at["X5", "f_accruals"] == "echec" and "Accruals élevés (6,0 %" in fb.at["X5", "motif"]
         and fb.loc["X0":"X4", "f_accruals"].eq("ok").all(), fb.at["X5", "motif"])

# Portefeuille : plafond sectoriel, une seule classe d'actions par société
p = pd.DataFrame({"signal": ["ACHETER"] * 8 + ["VENDRE", "GARDER"],
                  "score": [99, 98, 97, 96, 95, 94, 93, 92, 10, 70],
                  "secteur": ["A", "A", "A", "B", "B", "C", "D", "D", "A", "B"],
                  "cik": [1, 2, 3, 4, 4, 5, 6, 7, 8, 9], "motif": [""] * 10},
                 index=[f"P{i}" for i in range(10)])
cfg_p = sc.Config(n_positions=5, max_par_secteur=0.4)   # 2 lignes max par secteur
gardes, achats, ventes = sc.construire_portefeuille(p, ["P8", "P9"], cfg_p)
verifier("position GARDER conservée, VENDRE vendue", gardes == ["P9"] and [v[0] for v in ventes] == ["P8"])
verifier("plafond sectoriel et dédoublonnage par CIK respectés", achats == ["P0", "P1", "P3", "P5"],
         f"achats = {achats}")

# Métriques
jours = pd.bdate_range("2020-01-01", periods=505)
nav = pd.Series(np.exp(np.linspace(0, np.log(2), len(jours))), index=jours)
m = sc.metriques(nav, nav, None, est_ref=True)
attendu = 2 ** (365.25 / (jours[-1] - jours[0]).days) - 1
verifier("rendement annualisé exact", abs(m["Rendement annualisé"] - attendu) < 1e-9)
nav2 = nav.copy()
nav2.iloc[200:] = nav2.iloc[200:] * 0.7
attendu_mdd = 0.7 * nav.iloc[200] / nav.iloc[199] - 1
verifier("perte maximale exacte", abs(sc.metriques(nav2, nav, None)["Perte maximale"] - attendu_mdd) < 1e-12,
         f"{attendu_mdd:.5f}")

# Composition historique et renommages
sc.RENOMMAGES["OLD1"] = ("NEW1", None)
compo = sc.Composition(D["hist"], D["actuel"])
m16 = compo.membres("2016-01-01")
verifier("ancien ticker renommé suivi sous son nom actuel", "NEW1" in m16 and "OLD1" not in m16)
verifier("membre disparu présent dans la composition historique", "GONE" in compo.membres("2012-01-01")
         and "GONE" not in compo.membres("2016-01-01"))
verifier("titre ajouté en 2020 absent avant", "ADD1" not in compo.membres("2019-12-31")
         and "ADD1" in compo.membres("2020-06-01"))

# Divisions d'actions : capitalisation continue et F-score non pénalisé (société T07, 4 pour 1 en 2020)
t7 = sc.table_fondamentale(sc.extraire_faits(D["jsons"][1007]))
autour = t7[(t7["fin"] >= "2020-09-30") & (t7["fin"] <= "2021-06-30")]
verifier("division d'actions : pas comptée comme une émission (signal F-score neutralisé)",
         (autour["fscore_n"] == 8).all() and len(autour) >= 3, f"{autour['fscore_n'].tolist()}")
prix_t = {"adj": D["close"], "close": D["close"], "splits": D["splits"]}
close_t, spy_t, split_t, fact_t = sc._preparer_prix(prix_t)
mats_t = sc.matrices_prix(close_t, spy_t)
dates_t = [d for d in sc._fins_de_mois(close_t.index) if pd.Timestamp("2019-06-01") <= d <= pd.Timestamp("2022-06-30")]
pos_t = [close_t.index.get_loc(d) for d in dates_t]
mdd_t = sc.perte_max_1an(close_t, pos_t)
pan_t = sc.panneau_fondamental({1007: t7})
fond_t = sc.fondamentaux_aux_dates(pan_t, dates_t, [1007])
infos_t = D["actuel"].set_index("ticker")[["nom", "secteur", "cik"]]
impl = []
for d, i in zip(dates_t, pos_t):
    x = sc.tableau_facteurs(d, i, ["T07"], close_t, mats_t, mdd_t.loc[d], infos_t,
                            fond_t[fond_t["date"] == d], None, "complet", close_split=split_t, facteur=fact_t)
    impl.append(x.at["T07", "capitalisation"] / x.at["T07", "prix_split"])
impl = np.array(impl)
verifier("division d'actions : nombre d'actions implicite stable (pas de saut x4)",
         np.nanmax(impl) / np.nanmin(impl) < 1.15, f"min {np.nanmin(impl):.3g}, max {np.nanmax(impl):.3g}")

# Sociétés passées au XBRL en juillet 2010 : notées dès leur rapport annuel 2010 (publié début 2011)
t13 = sc.table_fondamentale(sc.extraire_faits(D["jsons"][1013]))
fy10 = t13[t13["annuel"] & (t13["fin"].dt.year == 2010)].iloc[0]
verifier("XBRL depuis 2010 : F-score calculable dès le rapport annuel 2010",
         fy10["fscore_n"] >= 6 and fy10["dispo"] < pd.Timestamp("2011-04-01"),
         f"{int(fy10['fscore_n'])} signaux, disponible le {fy10['dispo'].date()}")
verifier("XBRL depuis 2010 : surprise de résultats dès le premier trimestre publié",
         np.isfinite(t13["surprise_rn"].iloc[0]), f"premier trimestre {t13['fin'].iloc[0].date()}")
ttm_err = []
tri13, q13 = D["verites"][1013]["tri"], D["verites"][1013]["q"]
fins13 = [pd.Timestamp(r[3]) for r in tri13]
for _, ligne in t13.iterrows():
    if np.isfinite(ligne["resultat_net"]):
        k = fins13.index(ligne["fin"])
        ttm_err.append(abs(ligne["resultat_net"] / q13["rn"][k - 3:k + 1].sum() - 1))
verifier("XBRL depuis 2010 : TTM toujours exacts", max(ttm_err) < 1e-9, f"{len(ttm_err)} lignes")

# ---------------------------------------------------------------------------
print("\n3. Test complet sur données simulées (screener du jour + backtest + rapports)")
if os.path.exists(SORTIE):
    shutil.rmtree(SORTIE)


class FauxSEC:
    def __init__(self, cfg):
        if not cfg.email_valide():
            raise ValueError("email")
        self.appels = 0

    def table_tickers(self):
        t = {sc.cle_ticker(r.ticker): (int(r.cik), r.nom) for r in D["actuel"].itertuples()}
        return t

    def faits(self, cik):
        self.appels += 1
        return sc.extraire_faits(D["jsons"][int(cik)]) if int(cik) in D["jsons"] else pd.DataFrame(
            columns=sc.COLONNES_FAITS)

    def sic(self, cik):
        return 3570


def faux_prix(tickers, debut, cfg):
    c = D["close"]
    cols = [t for t in sorted(set(tickers)) if t in c.columns]
    m = c.index >= pd.Timestamp(debut)
    return {"adj": c.loc[m, cols], "close": c.loc[m, cols], "splits": D["splits"].loc[m, cols]}


sc.charger_sp500_actuel = lambda cfg: D["actuel"]
sc.charger_historique = lambda cfg: D["hist"]
sc.telecharger_prix = faux_prix
sc.ClientSEC = FauxSEC
sc.lire_revisions = lambda tickers, cfg: pd.Series(np.random.default_rng(3).normal(0, 0.05, len(tickers)),
                                                   index=list(tickers))
sc.lire_capitalisations = lambda tickers: pd.Series(2e10, index=list(tickers))
sc.charger_taux_sans_risque = lambda cfg: None      # FRED : réseau indisponible ici
sc.aujourdhui = lambda: pd.Timestamp("2026-10-16")
sc.RENOMMAGES["OLD1"] = ("NEW1", None)

cfg = sc.Config(email="test@domaine-test.org", dossier=SORTIE, n_positions=8)
t0 = time.time()
res = sc.lancer_screener(cfg, positions=["T00", "T05", "T20", "AAPL"])
dfc = res["classement"]
verifier("screener : tous les titres classés", len(dfc) == len(D["actuel"]), f"{len(dfc)} titres")
verifier("screener : scores entre 0 et 100", dfc["score"].dropna().between(0, 100).all())
verifier("screener : des signaux ACHETER existent", (dfc["signal"] == "ACHETER").sum() > 0,
         dict(dfc["signal"].value_counts()))
verifier("screener : le portefeuille ne contient que des titres valables",
         res["portefeuille"]["signal"].isin(["ACHETER", "GARDER", "SURVEILLER", "DONNÉES INSUFFISANTES"]).all())
verifier("screener : titre hors indice signalé, pas vendu",
         res["mes_positions"].set_index("ticker").at["AAPL", "action"].startswith("Non analysé"))
verifier("screener : valorisation calculée", dfc["rend_fcf"].notna().mean() > 0.6,
         f"{dfc['rend_fcf'].notna().mean():.0%} des titres")
verifier("screener : F-score calculé", dfc["fscore_9"].notna().mean() > 0.9)
k12 = len(D["verites"][1012]["actions"]) - 1
verifier("société multi-classes : capitalisation via le total d'actions du bilan",
         np.isfinite(dfc.at["T12", "capitalisation"]) and dfc.at["T12", "actions"] > 0,
         f"{dfc.at['T12', 'capitalisation']:.3g}")
achat = dfc[dfc["signal"] == "ACHETER"]
verifier("ACHETER respecte les seuils", (achat["score"] >= cfg.seuil_achat).all()
         and (achat["robustesse_pct"] >= cfg.robustesse_min_achat).all() and not achat["echec_filtre"].any())
print(f"  (screener simulé en {time.time() - t0:.1f} s)")

t0 = time.time()
bt = sc.lancer_backtest(cfg)
navs = bt["navs"]
verifier("backtest : séries complètes", navs.notna().all().all() and len(navs) > 3000, f"{len(navs)} jours")
verifier("backtest : valeur toujours positive", (navs["Stratégie"] > 0).all())
jr0 = bt["journal"]
verifier("backtest : démarre en janvier 2011", jr0["date"].iloc[0] == pd.Timestamp("2011-01-31"),
         str(jr0["date"].iloc[0].date()))
m11 = jr0[(jr0["date"] >= "2011-03-31") & (jr0["date"] <= "2011-12-31")]
part = (m11["insuffisantes"] / m11["titres"]).max()
verifier("backtest 2011 : moins de 15 % de titres sans données à partir de mars 2011", part < 0.15,
         f"maximum {part:.0%}")
reb11 = jr0[(jr0["date"] <= "2011-12-31") & jr0["rebalancement"]]["positions"].tolist()
verifier("backtest 2011 : positions dès le 1er mois, portefeuille rempli à 75 % au 2e rebalancement",
         reb11[0] >= 1 and reb11[1] >= 0.75 * cfg.n_positions, f"{reb11} lignes (univers simulé de 42 titres)")
cv = bt["couverture"]
verifier("backtest : couverture < 100 % avant 2015 (GONE sans cours)", cv[cv.index < "2015-06-01"].max() < 1.0,
         f"{cv[cv.index < '2015-06-01'].max():.3f}")
jr = bt["journal"]
verifier("backtest : nombre de lignes respecté", jr["positions"].max() <= cfg.n_positions)
an = bt["annuel"]
verifier("backtest : rendements annuels cohérents avec la courbe",
         abs((1 + an["Stratégie"]).prod() - navs["Stratégie"].iloc[-1]) < 1e-6)
ic = bt["analyse"]["ic"].set_index("Facteur")
verifier("analyse : l'IC de la qualité est positif (signal simulé lié à la qualité)",
         ic.at["Qualité", "IC moyen à 12 mois"] > 0, f"IC 12 mois = {ic.at['Qualité', 'IC moyen à 12 mois']:.3f}")
verifier("analyse : le faible risque est positif", ic.at["Faible risque", "IC moyen à 12 mois"] > 0,
         f"IC 12 mois = {ic.at['Faible risque', 'IC moyen à 12 mois']:.3f}")
print(f"  (backtest simulé en {time.time() - t0:.1f} s)")
for nom, tab in bt["metriques"].items():
    print(f"  {nom} :\n" + tab.round(3).to_string().replace("\n", "\n    "))

chemins = sc.generer_rapports(cfg, res, bt)
verifier("rapport Excel créé", os.path.getsize(chemins["excel"]) > 20000)
verifier("rapport HTML créé", os.path.getsize(chemins["html"]) > 50000)
import openpyxl  # noqa: E402
wb = openpyxl.load_workbook(chemins["excel"])
verifier("Excel : feuilles attendues", {"Classement", "Portefeuille proposé", "Mes positions", "BT Annuel",
                                         "BT Journal", "Paramètres"} <= set(wb.sheetnames), wb.sheetnames)
verifier("Excel : l'e-mail n'est pas exporté",
         all("domaine-test" not in str(c.value) for row in wb["Paramètres"].iter_rows() for c in row))

# Mode prix seul
cfg_p = sc.Config(email="test@domaine-test.org", dossier=SORTIE, mode="prix", debut_backtest="2010-06-01",
                  n_positions=8, revisions_analystes=False)
res_p = sc.lancer_screener(cfg_p)
bt_p = sc.lancer_backtest(cfg_p)
verifier("mode prix : screener et backtest fonctionnent", res_p["classement"]["score"].notna().mean() > 0.9
         and bt_p["navs"]["Stratégie"].notna().all())

# Séance en cours et raisons précises des données insuffisantes
idx = pd.DatetimeIndex(["2026-10-16", "2026-10-19"])
ny = lambda s: pd.Timestamp(s, tz="America/New_York")  # noqa: E731
verifier("cours du jour ignoré tant que la séance de New York est ouverte",
         sc.seance_en_cours(idx, ny("2026-10-19 12:20")) and not sc.seance_en_cours(idx, ny("2026-10-19 17:00"))
         and not sc.seance_en_cours(idx, ny("2026-10-20 09:00")))
jours_p = pd.bdate_range("2026-10-12", "2026-10-19")
px = pd.DataFrame({"SPY": 1.0, "AAA": 2.0}, index=jours_p)
vrai_seance = sc.seance_en_cours
sc.seance_en_cours = lambda index, maintenant=None: True          # séance du 19 encore ouverte
sc.ETAT_SOURCES["cours"] = {"derniere_seance": "2026-10-19"}
adj_p, spy_p, _, _ = sc._preparer_prix({"adj": px, "close": px, "splits": px * 0})
sc.seance_en_cours = vrai_seance
verifier("séance ouverte ignorée : la dernière séance affichée est la précédente, close",
         spy_p.index[-1] == pd.Timestamp("2026-10-16") and sc.ETAT_SOURCES["cours"]["derniere_seance"] == "2026-10-16")
raisons = [sc._raison_insuffisance(pd.Series(r), cfg) for r in [
    {"historique_ok": False}, {"historique_ok": True, "fin": pd.NaT},
    {"historique_ok": True, "fin": pd.Timestamp("2026-06-30"), "fscore_n": 4.0, "actif": 1e9},
    {"historique_ok": True, "fin": pd.Timestamp("2026-06-30"), "fscore_n": 8.0, "actif": float("nan")},
    {"historique_ok": True, "fin": pd.Timestamp("2026-06-30"), "fscore_n": 1.0, "actif": 1e9},
    {"historique_ok": True, "fin": pd.Timestamp("2026-06-30"), "fscore_n": 0.0, "actif": 1e9},
    {"historique_ok": True, "fin": pd.Timestamp("2026-06-30"), "fscore_n": 7.0, "actif": 1e9, "mom": float("nan")}]]
verifier("données insuffisantes : la raison précise est donnée", raisons[0].startswith("Moins de 200 séances")
         and raisons[1].startswith("Aucun compte récent") and "4 critères" in raisons[2] and "actif" in raisons[3]
         and "sur 1 critère seulement" in raisons[4] and raisons[5].startswith("F-score non mesurable")
         and raisons[6].startswith("Moins d'un an de cours"),
         " | ".join(raisons[2:]))

print(f"\nRésultat : {len(OK)} tests réussis, {len(KO)} échecs")
if KO:
    print("Échecs :", KO)
    sys.exit(1)
