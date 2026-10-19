#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
SCREENER S&P 500 — version 100 % gratuite
==========================================

Ce que fait ce programme
------------------------
1. Il récupère la composition du S&P 500, les cours ajustés (dividendes inclus)
   et les comptes déposés à la SEC.
2. Il élimine les « mines » avec des filtres de sécurité : F-score de Piotroski,
   couverture des intérêts, endettement, free cash-flow, qualité comptable
   (accruals) et volatilité extrême. Le M-score de Beneish est affiché en alerte.
3. Il note chaque action de 0 à 100 en la comparant aux titres de son secteur :
     - Potentiel  = momentum 12-1 mois + dynamique des résultats + valorisation
     - Robustesse = qualité (rentabilité, stabilité, endettement) + faible risque
4. Il en déduit un signal ACHETER / GARDER / SURVEILLER / VENDRE et propose un
   portefeuille de 25 lignes (25 % maximum par secteur).
5. En option, il rejoue la stratégie mois par mois depuis 2011 sur la
   composition historique de l'indice, frais compris, et la compare au S&P 500.

Sources gratuites
-----------------
- Composition actuelle : Wikipédia (secours : github.com/fja05680/sp500)
- Composition historique : github.com/fja05680/sp500
- Cours : Yahoo Finance via la librairie yfinance (non officielle, parfois limitée)
- Comptes : API XBRL officielle de la SEC (data.sec.gov). Les chiffres sont pris
  « tels que publiés la première fois » et datés par leur date de dépôt, donc le
  backtest n'utilise jamais une information qui n'était pas encore connue.
- Révisions des analystes (classement du jour uniquement) : Yahoo Finance

Limites à connaître
-------------------
- Yahoo ne fournit pas l'historique des sociétés disparues (rachats, faillites) :
  le backtest ne peut pas les détenir, ce qui l'embellit un peu. Compare-le
  surtout à « l'univers équipondéré », construit avec exactement les mêmes données.
- Les comptes XBRL de la SEC ne remontent qu'à 2009-2011 : backtest complet à
  partir de 2011, avec des comptes un peu plus minces la première année (les
  sociétés passées au XBRL en 2010 ne sont notées qu'à partir de leur rapport
  annuel 2010, publié début 2011). Le mode « prix » (sans comptes) remonte à 2001.
- Il n'existe pas d'historique gratuit des révisions d'analystes : elles ne
  servent qu'au classement du jour.

Avertissement : outil d'aide à la décision, pas un conseil en investissement.
"""
from __future__ import annotations

import argparse
import base64
import glob
import hashlib
import html
import io
import json
import logging
import math
import os
import re
import sys
import time
import warnings
from collections import Counter
from dataclasses import asdict, dataclass, field, fields
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
import requests

warnings.filterwarnings("ignore", category=FutureWarning)
warnings.filterwarnings("ignore", category=RuntimeWarning)

VERSION = "1.3.1"


# ===========================================================================
# 0. Paramètres
# ===========================================================================

@dataclass
class Config:
    """Tous les réglages du screener. Les valeurs par défaut suivent la méthode décrite."""

    email: str = "ton.email@exemple.com"   # OBLIGATOIRE pour la SEC : ton adresse de contact
    dossier: str = "screener_donnees"       # cache et résultats
    mode: str = "complet"                   # "complet" (cours + comptes) ou "prix" (cours seulement)

    # --- Portefeuille et signaux ---
    n_positions: int = 25
    max_par_secteur: float = 0.25           # 25 % des lignes au maximum par secteur
    seuil_achat: float = 90.0               # ACHETER : score >= 90 (top 10 %)
    seuil_garder: float = 60.0              # GARDER : score >= 60 (top 40 %)
    seuil_vente: float = 50.0               # VENDRE : score < 50 (sous la médiane)
    robustesse_min_achat: float = 50.0      # ACHETER exige aussi une robustesse >= 50
    poids_potentiel: float = 0.5            # score = 50 % potentiel + 50 % robustesse

    # --- Filtres de sécurité ---
    fscore_min: float = 6.0
    couverture_min: float = 3.0             # EBIT / charges d'intérêts
    couverture_min_secteurs: Dict[str, float] = field(
        default_factory=lambda: {"Utilities": 1.5, "Real Estate": 1.5})
    exclusion_levier_secteur: float = 0.10  # 10 % les plus endettés de chaque secteur
    exclusion_accruals_secteur: float = 0.10
    exclusion_volatilite: float = 0.10      # 10 % les plus volatils de l'indice
    secteurs_sans_filtre_fcf: Tuple[str, ...] = ("Utilities",)
    beneish_exclut: bool = False            # True = exclure aussi les M-score > -1,78

    # --- Backtest ---
    debut_backtest: str = "2011-01-01"
    fin_backtest: Optional[str] = None
    debut_validation: str = "2019-01-01"    # début de la période « hors échantillon »
    frais: float = 0.001                    # 0,10 % par achat ou vente
    rebalancement: str = "trimestriel"      # "trimestriel" ou "mensuel"
    ventes_urgence: bool = True             # vendre dès qu'un filtre casse, sans attendre

    # --- Données ---
    revisions_analystes: bool = True        # révisions BPA Yahoo (classement du jour)
    pause_sec: float = 0.12                 # la SEC autorise 10 requêtes par seconde
    cache_sec_jours: float = 3              # les comptes changent au rythme des publications
    cache_prix_jours: float = 1

    def chemin(self, *morceaux: str) -> str:
        p = os.path.join(self.dossier, *morceaux)
        os.makedirs(os.path.dirname(p) or ".", exist_ok=True)
        return p

    def email_valide(self) -> bool:
        e = (self.email or "").strip().lower()
        return bool(re.match(r"^[^@\s]+@[^@\s]+\.[a-z]{2,}$", e)) and "exemple" not in e

    # Réglages modifiables depuis un fichier (application web) : jamais l'e-mail ni le dossier.
    NON_MODIFIABLES = ("email", "dossier")

    @classmethod
    def depuis_dict(cls, d: Optional[dict], **forces) -> "Config":
        """Construit les réglages à partir d'un dictionnaire (fichier JSON), en vérifiant chaque valeur."""
        types = {f.name: str(f.type) for f in fields(cls)}
        valeurs = {}
        for k, v in (d or {}).items():
            if str(k).startswith("_") or k in cls.NON_MODIFIABLES:
                continue                                    # commentaires, ou réglages réservés
            if k not in types:
                raise ValueError(f"Réglage inconnu : « {k} ». Vérifie l'orthographe dans le fichier de réglages.")
            t = types[k]
            try:
                if t == "bool":
                    v = v if isinstance(v, bool) else str(v).strip().lower() in ("true", "vrai", "oui", "1")
                elif t == "int":
                    if isinstance(v, bool) or float(v) != int(float(v)):
                        raise ValueError
                    v = int(float(v))
                elif t == "float":
                    if isinstance(v, bool):
                        raise ValueError
                    v = float(v)
                elif t.startswith("Tuple"):
                    v = tuple(str(x) for x in ([v] if isinstance(v, str) else v))
                elif t.startswith("Dict"):
                    v = {str(a): float(b) for a, b in dict(v).items()}
                elif v is not None:
                    v = str(v)
            except (TypeError, ValueError):
                raise ValueError(f"Réglage « {k} » : valeur invalide ({v!r}).") from None
            valeurs[k] = v
        valeurs.update(forces)
        cfg = cls(**valeurs)
        cfg.valider()
        return cfg

    def valider(self) -> None:
        def exiger(cond: bool, texte: str):
            if not cond:
                raise ValueError(f"Réglage invalide : {texte}")
        exiger(self.mode in ("complet", "prix"), "mode doit valoir « complet » ou « prix »")
        exiger(self.rebalancement in ("trimestriel", "mensuel"), "rebalancement : « trimestriel » ou « mensuel »")
        exiger(3 <= self.n_positions <= 100, "n_positions entre 3 et 100")
        exiger(0 < self.max_par_secteur <= 1, "max_par_secteur entre 0 et 1 (0.25 = 25 %)")
        exiger(0 <= self.seuil_vente <= self.seuil_garder <= self.seuil_achat <= 100,
               "il faut 0 ≤ seuil_vente ≤ seuil_garder ≤ seuil_achat ≤ 100")
        exiger(0 <= self.robustesse_min_achat <= 100, "robustesse_min_achat entre 0 et 100")
        exiger(0 <= self.poids_potentiel <= 1, "poids_potentiel entre 0 et 1")
        exiger(0 <= self.fscore_min <= 9, "fscore_min entre 0 et 9")
        for nom in ("exclusion_levier_secteur", "exclusion_accruals_secteur", "exclusion_volatilite"):
            exiger(0 <= getattr(self, nom) < 1, f"{nom} entre 0 et 1 (0.10 = 10 %)")
        exiger(0 <= self.frais < 0.05, "frais entre 0 et 0.05 (0.001 = 0,10 %)")
        for nom in ("debut_backtest", "debut_validation"):
            try:
                pd.Timestamp(getattr(self, nom))
            except (TypeError, ValueError):
                exiger(False, f"{nom} doit être une date AAAA-MM-JJ")


# ===========================================================================
# 1. Utilitaires
# ===========================================================================

def info(msg: str) -> None:
    print(f"[{time.strftime('%H:%M:%S')}] {msg}", flush=True)


def aujourdhui() -> pd.Timestamp:
    return pd.Timestamp.today().normalize()


def _frais(chemin: str, jours: float) -> bool:
    return os.path.exists(chemin) and (time.time() - os.path.getmtime(chemin)) < jours * 86400


def ticker_yahoo(sym) -> str:
    s = str(sym).strip().upper()
    return s if s.startswith("^") else s.replace(".", "-").replace("/", "-")


def cle_ticker(sym) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(sym).upper())


def http_get(url: str, entetes: Optional[dict] = None, timeout: int = 45, essais: int = 4) -> requests.Response:
    derniere = None
    for k in range(essais):
        try:
            r = requests.get(url, headers=entetes or {}, timeout=timeout)
        except requests.RequestException as e:
            derniere = e
            time.sleep(2.0 * (k + 1))
            continue
        if r.status_code == 200:
            return r
        if r.status_code in (403, 429, 500, 502, 503, 504):
            derniere = f"HTTP {r.status_code}"
            time.sleep(3.0 * (k + 1) ** 2)
            continue
        raise RuntimeError(f"HTTP {r.status_code} pour {url}")
    raise RuntimeError(f"Échec après {essais} essais : {url} ({derniere})")


def _ua_web(cfg: Config) -> dict:
    # L'adresse e-mail n'est envoyée qu'à la SEC, jamais à Wikipédia ni à GitHub. Comme contact, l'adresse
    # publique du dépôt quand le programme tourne sur GitHub Actions (règle d'usage de Wikipédia).
    depot = os.environ.get("GITHUB_REPOSITORY", "").strip()
    contact = f"+https://github.com/{depot}" if re.match(r"^[\w.-]+/[\w.-]+$", depot) else "outil personnel d'analyse"
    return {"User-Agent": f"ScreenerSP500/{VERSION} ({contact}) python-requests/{requests.__version__}"}


SECTEURS_FR = {
    "Information Technology": "Technologies de l'information",
    "Health Care": "Santé",
    "Financials": "Finance",
    "Consumer Discretionary": "Consommation discrétionnaire",
    "Communication Services": "Services de communication",
    "Industrials": "Industrie",
    "Consumer Staples": "Consommation de base",
    "Energy": "Énergie",
    "Utilities": "Services aux collectivités",
    "Real Estate": "Immobilier",
    "Materials": "Matériaux",
    "Inconnu": "Inconnu",
}


def secteur_depuis_sic(sic) -> str:
    """Correspondance approximative code SIC (SEC) -> secteur GICS, pour les anciens membres."""
    try:
        s = int(sic)
    except (TypeError, ValueError):
        return "Inconnu"
    if s <= 0:
        return "Inconnu"
    regles = [
        (100, 999, "Consumer Staples"), (1000, 1099, "Materials"), (1200, 1399, "Energy"),
        (1400, 1499, "Materials"), (1531, 1531, "Consumer Discretionary"), (1500, 1799, "Industrials"),
        (2000, 2199, "Consumer Staples"), (2200, 2399, "Consumer Discretionary"), (2400, 2499, "Materials"),
        (2500, 2599, "Consumer Discretionary"), (2600, 2699, "Materials"), (2700, 2799, "Communication Services"),
        (2830, 2836, "Health Care"), (2840, 2844, "Consumer Staples"), (2800, 2899, "Materials"),
        (2900, 2999, "Energy"), (3011, 3011, "Consumer Discretionary"), (3000, 3099, "Materials"),
        (3100, 3199, "Consumer Discretionary"), (3200, 3399, "Materials"), (3400, 3499, "Industrials"),
        (3570, 3579, "Information Technology"), (3500, 3599, "Industrials"),
        (3630, 3639, "Consumer Discretionary"), (3651, 3651, "Consumer Discretionary"),
        (3660, 3679, "Information Technology"), (3600, 3699, "Industrials"),
        (3710, 3716, "Consumer Discretionary"), (3750, 3751, "Consumer Discretionary"), (3700, 3799, "Industrials"),
        (3841, 3851, "Health Care"), (3820, 3829, "Information Technology"), (3800, 3899, "Industrials"),
        (3900, 3999, "Consumer Discretionary"), (4000, 4799, "Industrials"), (4800, 4899, "Communication Services"),
        (4950, 4959, "Industrials"), (4900, 4999, "Utilities"), (5122, 5122, "Health Care"),
        (5140, 5149, "Consumer Staples"), (5180, 5182, "Consumer Staples"), (5000, 5199, "Industrials"),
        (5331, 5331, "Consumer Staples"), (5399, 5399, "Consumer Staples"), (5400, 5499, "Consumer Staples"),
        (5912, 5912, "Consumer Staples"), (5200, 5999, "Consumer Discretionary"), (6500, 6599, "Real Estate"),
        (6798, 6798, "Real Estate"), (6000, 6799, "Financials"), (7370, 7379, "Information Technology"),
        (7800, 7899, "Communication Services"), (7000, 7099, "Consumer Discretionary"),
        (7200, 7299, "Consumer Discretionary"), (7900, 7999, "Consumer Discretionary"), (7000, 7999, "Industrials"),
        (8000, 8099, "Health Care"), (8100, 8999, "Industrials"),
    ]
    for a, b, secteur in regles:
        if a <= s <= b:
            return secteur
    return "Industrials"


# ===========================================================================
# 2. Univers : composition actuelle et historique du S&P 500
# ===========================================================================

URL_WIKI = "https://en.wikipedia.org/wiki/List_of_S%26P_500_companies"
URL_GITHUB = "https://raw.githubusercontent.com/fja05680/sp500/master/"
FICHIER_HISTORIQUE = "S%26P%20500%20Historical%20Components%20%26%20Changes%20(Updated).csv"

# Anciens tickers d'entreprises qui existent toujours sous un autre nom chez Yahoo.
# ancien ticker -> (ticker Yahoo actuel, à appliquer avant cette date, ou None = toujours)
RENOMMAGES: Dict[str, Tuple[str, Optional[str]]] = {
    "FB": ("META", None), "ANTM": ("ELV", None), "FI": ("FISV", None), "RE": ("EG", None),
    "ABC": ("COR", None), "PKI": ("RVTY", None), "FLT": ("CPAY", None), "WLTW": ("WTW", None),
    "HCP": ("DOC", None), "PEAK": ("DOC", None), "DISCA": ("WBD", None), "COG": ("CTRA", None),
    "MYL": ("VTRS", None), "UTX": ("RTX", None), "KORS": ("CPRI", None), "CTL": ("LUMN", None),
    "BBT": ("TFC", None), "HRS": ("LHX", None), "JEC": ("J", None), "DWDP": ("DD", None),
    "BLL": ("BALL", None), "SYMC": ("GEN", None), "NLOK": ("GEN", None), "CDAY": ("DAY", None),
    "ADS": ("BFH", None), "BHGE": ("BKR", None), "LB": ("BBWI", None),
    "ARNC": ("HWM", "2020-04-06"),   # Arconic (ancien) = Howmet aujourd'hui
    "IR": ("TT", "2020-03-03"),      # Ingersoll-Rand plc = Trane Technologies aujourd'hui
}

# Sociétés dont le nombre d'actions SEC n'est pas exploitable (classes A/B très différentes)
CAPI_NON_FIABLE_SEC = {"BRK-B"}


def charger_sp500_actuel(cfg: Config) -> pd.DataFrame:
    """Composition actuelle : ticker, nom, secteur GICS, sous-secteur, CIK."""
    chemin = cfg.chemin("cache", "sp500_actuel.pkl")
    if _frais(chemin, 1):
        return pd.read_pickle(chemin)
    brut, source = None, ""
    try:
        page = http_get(URL_WIKI, _ua_web(cfg)).text
        for t in pd.read_html(io.StringIO(page)):
            cols = [str(c) for c in t.columns]
            if "Symbol" in cols and "GICS Sector" in cols:
                brut, source = t, "Wikipédia"
                break
    except Exception as e:  # noqa: BLE001
        info(f"Wikipédia indisponible ({type(e).__name__}) : utilisation de la copie GitHub.")
    if brut is None:
        texte = http_get(URL_GITHUB + "sp500.csv", _ua_web(cfg)).text
        brut, source = pd.read_csv(io.StringIO(texte)), "GitHub fja05680/sp500"
    n = len(brut)
    out = pd.DataFrame({
        "ticker": brut["Symbol"].map(ticker_yahoo).values,
        "nom": brut["Security"].astype(str).values,
        "secteur": brut["GICS Sector"].astype(str).values,
        "sous_secteur": (brut["GICS Sub-Industry"].astype(str).values
                         if "GICS Sub-Industry" in brut else [""] * n),
        "cik": (pd.to_numeric(brut["CIK"], errors="coerce").values if "CIK" in brut else [np.nan] * n),
    })
    out = out.drop_duplicates("ticker").reset_index(drop=True)
    out.attrs["source"] = source
    out.to_pickle(chemin)
    info(f"Composition actuelle : {len(out)} titres ({source}).")
    return out


def charger_historique(cfg: Config) -> pd.DataFrame:
    """Composition historique (une ligne par changement) : date, liste de tickers."""
    chemin = cfg.chemin("cache", "sp500_historique.pkl")
    if _frais(chemin, 7):
        return pd.read_pickle(chemin)
    texte = http_get(URL_GITHUB + FICHIER_HISTORIQUE, _ua_web(cfg)).text
    h = pd.read_csv(io.StringIO(texte))
    h["date"] = pd.to_datetime(h["date"])
    h["tickers"] = h["tickers"].astype(str).map(
        lambda s: sorted({x.strip().upper() for x in s.split(",") if x.strip()}))
    h = h.sort_values("date").reset_index(drop=True)[["date", "tickers"]]
    h.to_pickle(chemin)
    info(f"Composition historique : {len(h)} dates, du {h['date'].iloc[0].date()} au {h['date'].iloc[-1].date()}.")
    return h


def vers_yahoo_historique(sym: str, date: pd.Timestamp) -> str:
    s = str(sym).strip().upper()
    if s in RENOMMAGES:
        nouveau, avant = RENOMMAGES[s]
        if avant is None or date < pd.Timestamp(avant):
            s = nouveau
    return ticker_yahoo(s)


class Composition:
    """Membres de l'indice à une date donnée (tickers au format Yahoo)."""

    def __init__(self, historique: pd.DataFrame, actuel: pd.DataFrame):
        self.dates = historique["date"].values.astype("datetime64[ns]")
        self.listes = historique["tickers"].tolist()
        self.actuels = set(actuel["ticker"])
        self.derniere = pd.Timestamp(historique["date"].max())

    def brut(self, date) -> List[str]:
        date = pd.Timestamp(date)
        if date > self.derniere:
            return sorted(self.actuels)
        i = int(np.searchsorted(self.dates, np.datetime64(date, "ns"), side="right")) - 1
        return list(self.listes[i]) if i >= 0 else []

    def membres(self, date) -> set:
        date = pd.Timestamp(date)
        if date > self.derniere:
            return set(self.actuels)
        return {vers_yahoo_historique(s, date) for s in self.brut(date)}

    def tous_depuis(self, debut, fin) -> set:
        debut, fin = pd.Timestamp(debut), pd.Timestamp(fin)
        res = set(self.membres(debut))
        for d, liste in zip(self.dates, self.listes):
            d = pd.Timestamp(d)
            if debut <= d <= fin:
                res |= {vers_yahoo_historique(s, d) for s in liste}
        if fin > self.derniere:
            res |= self.actuels
        return res


# ===========================================================================
# 3. Cours (Yahoo Finance via yfinance)
# ===========================================================================

class _Capture(logging.Handler):
    def __init__(self):
        super().__init__(level=logging.DEBUG)
        self.messages: List[str] = []

    def emit(self, record):
        try:
            self.messages.append(record.getMessage())
        except Exception:  # noqa: BLE001
            pass


def _extraire_champ(brut, lot: Sequence[str], champ: str) -> pd.DataFrame:
    if brut is None or len(brut) == 0:
        return pd.DataFrame()
    if isinstance(brut.columns, pd.MultiIndex):
        n0 = set(brut.columns.get_level_values(0))
        n1 = set(brut.columns.get_level_values(1))
        if champ in n0:
            c = brut[champ]
        elif champ in n1:
            c = brut.xs(champ, axis=1, level=1)
        else:
            return pd.DataFrame()
    else:
        if champ not in brut.columns:
            return pd.DataFrame()
        c = brut[[champ]].rename(columns={champ: lot[0]})
    if isinstance(c, pd.Series):
        c = c.to_frame(lot[0])
    c = c.apply(pd.to_numeric, errors="coerce")
    c.columns = [str(x) for x in c.columns]
    return c


CHAMPS_PRIX = ("adj", "close", "splits")


def _normaliser_index(df: pd.DataFrame) -> pd.DataFrame:
    idx = pd.to_datetime(df.index)
    if getattr(idx, "tz", None) is not None:
        idx = idx.tz_localize(None)
    df.index = idx.normalize()
    return df[~df.index.duplicated(keep="last")].sort_index()


def _yf_telecharger(tickers: Sequence[str], debut: pd.Timestamp, taille: int = 60) -> Dict[str, pd.DataFrame]:
    """Télécharge, par lots, la clôture ajustée (rendement total), la clôture ajustée des seules
    divisions d'actions (pour la capitalisation) et l'historique des divisions."""
    import yfinance as yf

    journal = logging.getLogger("yfinance")
    morceaux: Dict[str, List[pd.DataFrame]] = {k: [] for k in CHAMPS_PRIX}
    for i in range(0, len(tickers), taille):
        lot = list(tickers[i:i + taille])
        obtenus: Dict[str, Dict[str, pd.Series]] = {k: {} for k in CHAMPS_PRIX}
        for essai in range(4):
            restants = [t for t in lot if t not in obtenus["adj"]]
            if not restants:
                break
            capture = _Capture()
            anciens, propage = journal.handlers[:], journal.propagate
            journal.handlers, journal.propagate = [capture], False
            try:
                brut = yf.download(restants, start=debut.strftime("%Y-%m-%d"), auto_adjust=False,
                                   actions=True, progress=False, threads=True, group_by="column",
                                   timeout=30)
                adj = _extraire_champ(brut, restants, "Adj Close")
                close = _extraire_champ(brut, restants, "Close")
                splits = _extraire_champ(brut, restants, "Stock Splits")
                if adj.empty:
                    adj = close
            except Exception as e:  # noqa: BLE001
                capture.messages.append(f"{type(e).__name__}: {e}")
                adj = close = splits = pd.DataFrame()
            finally:
                journal.handlers, journal.propagate = anciens, propage
            for t in adj.columns:
                if adj[t].notna().any():
                    obtenus["adj"][t] = adj[t]
                    obtenus["close"][t] = close[t] if t in close.columns else adj[t]
                    obtenus["splits"][t] = splits[t] if t in splits.columns else adj[t] * 0.0
            texte = " ".join(capture.messages).lower()
            limite = any(k in texte for k in ("rate limit", "too many requests", "ratelimit"))
            if not limite or essai == 3:
                break
            attente = 45 * (essai + 1)
            info(f"  Yahoo limite les requêtes : pause de {attente} s puis nouvel essai…")
            time.sleep(attente)
        for k in CHAMPS_PRIX:
            if obtenus[k]:
                morceaux[k].append(pd.DataFrame({t: obtenus[k][t] for t in lot if t in obtenus[k]}))
        info(f"  cours : {min(i + taille, len(tickers))}/{len(tickers)} titres traités")
        time.sleep(1.0)
    out = {}
    for k in CHAMPS_PRIX:
        if morceaux[k]:
            df = pd.concat(morceaux[k], axis=1)
            out[k] = _normaliser_index(df.loc[:, ~df.columns.duplicated()]).astype(float)
        else:
            out[k] = pd.DataFrame()
    return out


def facteur_divisions(splits: pd.DataFrame, index: pd.DatetimeIndex) -> pd.DataFrame:
    """F(t) = produit des divisions d'actions survenues après t (1 si aucune).
    Capitalisation à t = clôture ajustée des divisions(t) × actions publiées × F(date des actions)."""
    s = splits.fillna(0.0)
    r = s.where(s > 0, 1.0)
    cumul = r.iloc[::-1].cumprod().iloc[::-1]
    f = cumul.shift(-1).fillna(1.0)
    return f.reindex(index, method="ffill").bfill().fillna(1.0)


# --- Source de secours facultative : API officielle d'Alpaca (compte gratuit, historique depuis 2016) ---
URL_ALPACA = "https://data.alpaca.markets/v2/stocks/bars"
ETAT_SOURCES: Dict[str, dict] = {}     # ce que chaque source a fourni (affiché par l'application web)


def cles_alpaca() -> Optional[Tuple[str, str]]:
    """Clés lues dans l'environnement (secrets GitHub), jamais dans les réglages ni les exports."""
    cle = os.environ.get("ALPACA_API_KEY_ID", "").strip()
    secret = os.environ.get("ALPACA_API_SECRET_KEY", "").strip()
    return (cle, secret) if cle and secret else None


def _alpaca_closes(symboles: Sequence[str], debut: pd.Timestamp, ajustement: str, cles: Tuple[str, str],
                   flux: str = "sip") -> Dict[str, pd.Series]:
    """Clôtures quotidiennes d'Alpaca pour plusieurs symboles (toutes les pages de résultats)."""
    entetes = {"APCA-API-KEY-ID": cles[0], "APCA-API-SECRET-KEY": cles[1], "Accept": "application/json"}
    points: Dict[str, list] = {}
    jeton, essais = None, 0
    for _ in range(5000):
        params = {"symbols": ",".join(symboles), "timeframe": "1Day", "start": debut.strftime("%Y-%m-%d"),
                  "adjustment": ajustement, "feed": flux, "limit": 10000, "sort": "asc"}
        if jeton:
            params["page_token"] = jeton
        r = requests.get(URL_ALPACA, headers=entetes, params=params, timeout=60)
        if r.status_code == 429 and essais < 5:          # 200 requêtes par minute au maximum
            essais += 1
            time.sleep(15 * essais)
            continue
        if r.status_code != 200:
            raise RuntimeError(f"Alpaca HTTP {r.status_code} : {r.text[:160]}")
        js = r.json()
        for sym, barres in (js.get("bars") or {}).items():
            points.setdefault(sym, []).extend((b["t"], b["c"]) for b in barres or [])
        jeton = js.get("next_page_token")
        if not jeton:
            break
        time.sleep(0.32)
    out = {}
    for sym, pts in points.items():
        if not pts:
            continue
        idx = pd.to_datetime([t for t, _ in pts], utc=True).tz_convert("America/New_York").tz_localize(None)
        out[sym] = pd.Series([float(c) for _, c in pts], index=idx.normalize())
    return out


def _alpaca_telecharger(tickers: Sequence[str], debut: pd.Timestamp, cles: Tuple[str, str],
                        taille: int = 50) -> Dict[str, pd.DataFrame]:
    """Même format que Yahoo : rendement total, cours ajusté des divisions, divisions d'actions."""
    vers_alp = {t: t.replace("-", ".") for t in tickers if not t.startswith("^")}
    vers_yah = {a: t for t, a in vers_alp.items()}
    debut = max(pd.Timestamp(debut), pd.Timestamp("2016-01-01"))
    morceaux = {k: {} for k in ("all", "split", "raw")}
    syms = list(vers_alp.values())
    for i in range(0, len(syms), taille):
        lot = syms[i:i + taille]
        for aj in morceaux:
            try:
                res = _alpaca_closes(lot, debut, aj, cles)
            except RuntimeError as e:
                if "403" not in str(e):
                    raise
                res = _alpaca_closes(lot, debut, aj, cles, flux="iex")   # abonnement sans données SIP
            morceaux[aj].update(res)
    if not morceaux["all"]:
        return {k: pd.DataFrame() for k in CHAMPS_PRIX}
    adj = pd.DataFrame({vers_yah[s]: v for s, v in morceaux["all"].items() if s in vers_yah})
    close = pd.DataFrame({vers_yah[s]: v for s, v in morceaux["split"].items() if s in vers_yah}).reindex_like(adj)
    brut = pd.DataFrame({vers_yah[s]: v for s, v in morceaux["raw"].items() if s in vers_yah}).reindex_like(adj)
    # facteur cumulé des divisions à venir = brut / ajusté des divisions ; une division = saut de ce facteur
    f = (brut / close).round(4).ffill().bfill()
    saut = (f.shift(1) / f).round(4)
    splits = saut.where(saut > 1.0001, 0.0).fillna(0.0)
    return {"adj": _normaliser_index(adj).astype(float), "close": _normaliser_index(close.fillna(adj)).astype(float),
            "splits": _normaliser_index(splits).astype(float)}


def _completer_alpaca(cache: Dict[str, pd.DataFrame], manquants: Sequence[str], debut: pd.Timestamp
                      ) -> Tuple[Dict[str, pd.DataFrame], List[str]]:
    cles = cles_alpaca()
    if not cles or not manquants:
        return cache, []
    info(f"Cours : {len(manquants)} titres absents de Yahoo, essai avec Alpaca (source de secours)…")
    try:
        neuf = _alpaca_telecharger(list(manquants), debut, cles)
    except Exception as e:  # noqa: BLE001
        info(f"  Alpaca indisponible ({type(e).__name__}: {str(e)[:120]}).")
        return cache, []
    if neuf["adj"].empty:
        return cache, []
    obtenus = [t for t in neuf["adj"].columns if neuf["adj"][t].notna().any()]
    if cache is None or cache["adj"].empty:
        return {k: neuf[k][obtenus] for k in CHAMPS_PRIX}, obtenus
    fusion = {}
    for k in CHAMPS_PRIX:
        a = neuf[k][obtenus].reindex(cache[k].index.union(neuf[k].index))
        fusion[k] = pd.concat([cache[k].reindex(a.index), a], axis=1)
    info(f"  Alpaca : {len(obtenus)} titres complétés.")
    return fusion, obtenus


def telecharger_prix(tickers: Sequence[str], debut, cfg: Config) -> Dict[str, pd.DataFrame]:
    """Cours avec cache du jour. Renvoie {'adj': rendement total, 'close': ajusté des seules divisions,
    'splits': divisions d'actions}. Source : Yahoo, complétée au besoin par Alpaca (si des clés sont fournies)."""
    tickers = sorted({ticker_yahoo(t) for t in tickers})
    debut = pd.Timestamp(debut).normalize()
    chemin = cfg.chemin("cache", "prix.pkl")
    chemin_meta = cfg.chemin("cache", "prix_meta.json")
    cache, meta = None, {}
    if os.path.exists(chemin) and os.path.exists(chemin_meta):
        try:
            cache = pd.read_pickle(chemin)
            with open(chemin_meta, encoding="utf-8") as fh:
                meta = json.load(fh)
            if not isinstance(cache, dict) or set(cache) != set(CHAMPS_PRIX):
                cache = None
        except Exception:  # noqa: BLE001
            cache, meta = None, {}
    valide = (cache is not None and _frais(chemin, cfg.cache_prix_jours)
              and pd.Timestamp(meta.get("debut", "2100-01-01")) <= debut)
    if valide:
        echecs = set(meta.get("echecs", []))
        manquants = [t for t in tickers if t not in cache["adj"].columns and t not in echecs]
        if manquants:
            info(f"Cours : {len(manquants)} titres à compléter…")
            neuf = _yf_telecharger(manquants, pd.Timestamp(meta["debut"]))
            if not neuf["adj"].empty:
                cache = {k: pd.concat([cache[k], neuf[k]], axis=1) for k in CHAMPS_PRIX}
            rates = [t for t in manquants if t not in cache["adj"].columns]
            cache, par_alpaca = _completer_alpaca(cache, rates, pd.Timestamp(meta["debut"]))
            meta["alpaca"] = sorted(set(meta.get("alpaca", [])) | set(par_alpaca))
            meta["echecs"] = sorted(echecs | (set(rates) - set(par_alpaca)))
    else:
        info(f"Cours : téléchargement de {len(tickers)} titres depuis le {debut.date()} "
             "(quelques minutes la première fois)…")
        cache = _yf_telecharger(tickers, debut)
        rates = [t for t in tickers if t not in cache["adj"].columns]
        cache, par_alpaca = _completer_alpaca(cache, rates, debut)
        meta = {"debut": str(debut.date()), "echecs": sorted(set(rates) - set(par_alpaca)),
                "alpaca": sorted(par_alpaca)}
    if cache is None or cache["adj"].empty:
        raise RuntimeError("Aucun cours n'a pu être téléchargé depuis Yahoo Finance. "
                           "Réessaie dans quelques minutes (limite de requêtes) ou mets yfinance à jour.")
    cache = {k: cache[k].loc[:, ~cache[k].columns.duplicated()].sort_index() for k in CHAMPS_PRIX}
    ETAT_SOURCES["cours"] = {"demandes": len(tickers),
                             "obtenus": int(sum(t in cache["adj"].columns for t in tickers)),
                             "par_alpaca": [t for t in meta.get("alpaca", []) if t in tickers],
                             "derniere_seance": str(cache["adj"].index.max().date())}
    pd.to_pickle(cache, chemin)
    with open(chemin_meta, "w", encoding="utf-8") as fh:
        json.dump(meta, fh)
    presents = [t for t in tickers if t in cache["adj"].columns]
    manque = len(tickers) - len(presents)
    if manque:
        info(f"Cours : {len(presents)} titres disponibles, {manque} sans données chez Yahoo "
             "(sociétés rachetées, en faillite ou ticker changé).")
    return {k: cache[k].reindex(columns=presents) for k in CHAMPS_PRIX}


# ===========================================================================
# 4. Comptes : API XBRL de la SEC
# ===========================================================================

TAGS: Dict[str, List[str]] = {
    # --- Compte de résultat et flux (valeurs sur une durée) ---
    "revenus": ["Revenues", "RevenueFromContractWithCustomerExcludingAssessedTax",
                "RevenueFromContractWithCustomerIncludingAssessedTax", "SalesRevenueNet",
                "SalesRevenueGoodsNet", "SalesRevenueServicesNet", "RevenuesNetOfInterestExpense",
                "RegulatedAndUnregulatedOperatingRevenue", "ElectricUtilityRevenue",
                "OperatingLeasesIncomeStatementLeaseRevenue", "RealEstateRevenueNet",
                "HealthCareOrganizationRevenue", "OilAndGasRevenue", "InterestAndDividendIncomeOperating"],
    "cout_revenus": ["CostOfRevenue", "CostOfGoodsAndServicesSold", "CostOfGoodsSold", "CostOfServices",
                     "CostOfGoodsAndServiceExcludingDepreciationDepletionAndAmortization"],
    "marge_brute": ["GrossProfit"],
    # Total des charges d'exploitation : CA = résultat d'exploitation + charges quand le CA n'a qu'une
    # étiquette propre à la société, absente de l'API (APA Corporation)
    "couts_totaux": ["CostsAndExpenses"],
    "resultat_exploitation": ["OperatingIncomeLoss"],
    "resultat_avant_impot": [
        "IncomeLossFromContinuingOperationsBeforeIncomeTaxesExtraordinaryItemsNoncontrollingInterest",
        "IncomeLossFromContinuingOperationsBeforeIncomeTaxesMinorityInterestAndIncomeLossFromEquityMethodInvestments"],
    "resultat_net": ["NetIncomeLoss", "NetIncomeLossAvailableToCommonStockholdersBasic", "ProfitLoss"],
    "impot": ["IncomeTaxExpenseBenefit"],
    "charges_interets": ["InterestExpense", "InterestExpenseNonoperating", "InterestExpenseDebt",
                         "InterestAndDebtExpense", "InterestExpenseOperating", "InterestExpenseLongTermDebt",
                         "InterestCostsIncurred"],      # intérêts encourus, y compris capitalisés (promoteurs)
    # Intérêts nets (négatifs = charge nette) et intérêts payés : solutions de repli pour la couverture
    "interets_nets": ["InterestIncomeExpenseNonoperatingNet", "InterestIncomeExpenseNet"],
    "interets_payes": ["InterestPaidNet", "InterestPaid"],
    # Émissions et remboursements d'emprunts : preuve qu'une dette existe même si le bilan ne la nomme pas
    "flux_dette": ["RepaymentsOfLongTermDebt", "ProceedsFromIssuanceOfLongTermDebt", "RepaymentsOfDebt",
                   "ProceedsFromIssuanceOfDebt", "RepaymentsOfSeniorDebt", "ProceedsFromIssuanceOfSeniorLongTermDebt",
                   "RepaymentsOfNotesPayable", "ProceedsFromNotesPayable",
                   "RepaymentsOfDebtAndCapitalLeaseObligations", "RepaymentsOfLongTermDebtAndCapitalSecurities"],
    "cfo": ["NetCashProvidedByUsedInOperatingActivities",
            "NetCashProvidedByUsedInOperatingActivitiesContinuingOperations"],
    "capex": ["PaymentsToAcquirePropertyPlantAndEquipment", "PaymentsToAcquireProductiveAssets",
              "PaymentsForCapitalImprovements", "PaymentsToAcquireOtherPropertyPlantAndEquipment",
              "PaymentsToAcquireOtherProductiveAssets", "PaymentsToAcquireOilAndGasPropertyAndEquipment",
              "PaymentsToAcquireOilAndGasProperty", "PaymentsToExploreAndDevelopOilAndGasProperties",
              "PaymentsToDevelopRealEstateAssets",
              "SegmentExpenditureAdditionToLongLivedAssets"],   # total des secteurs (ASU 2023-07), en dernier recours
    "amortissements": ["DepreciationDepletionAndAmortization", "DepreciationAmortizationAndAccretionNet",
                       "DepreciationAndAmortization", "Depreciation"],
    "sga": ["SellingGeneralAndAdministrativeExpense", "GeneralAndAdministrativeExpense"],
    # --- Bilan (valeurs à une date) ---
    "actif": ["Assets"],
    "actif_courant": ["AssetsCurrent"],
    "passif_courant": ["LiabilitiesCurrent"],
    "passif": ["Liabilities"],
    "passif_et_cp": ["LiabilitiesAndStockholdersEquity"],
    "capitaux_propres": ["StockholdersEquity",
                         "StockholdersEquityIncludingPortionAttributableToNoncontrollingInterest"],
    "tresorerie": ["CashAndCashEquivalentsAtCarryingValue",
                   "CashCashEquivalentsRestrictedCashAndRestrictedCashEquivalents", "CashAndDueFromBanks", "Cash"],
    "placements_ct": ["ShortTermInvestments", "MarketableSecuritiesCurrent",
                      "AvailableForSaleSecuritiesDebtSecuritiesCurrent"],
    # Dette : du plus détaillé au plus agrégé (voir composer() dans table_fondamentale)
    "dette_totale": ["DebtLongtermAndShorttermCombinedAmount"],                 # long + court terme
    "dette_lt_totale": ["LongTermDebt", "LongTermDebtAndCapitalLeaseObligationsIncludingCurrentMaturities",
                        "DebtAndCapitalLeaseObligations", "NotesPayable"],      # long terme, échéances < 1 an incluses
    "dette_lt_nc": ["LongTermDebtNoncurrent", "LongTermDebtAndCapitalLeaseObligations", "LongTermNotesPayable",
                    "SeniorLongTermNotes", "UnsecuredLongTermDebt", "ConvertibleDebtNoncurrent",
                    "ConvertibleLongTermNotesPayable", "OtherLongTermDebtNoncurrent"],   # long terme, part > 1 an
    "dette_lt_courante": ["LongTermDebtCurrent", "LongTermDebtAndCapitalLeaseObligationsCurrent",
                          "NotesPayableCurrent", "ConvertibleNotesPayableCurrent"],      # échéances < 1 an
    "dette_courante": ["DebtCurrent"],                                          # tout le court terme
    "emprunts_ct": ["ShortTermBorrowings", "CommercialPaper", "OtherShortTermBorrowings", "LinesOfCreditCurrent"],
    # Lignes du bilan des foncières (REIT), additionnées faute de total : obligations, emprunts non garantis,
    # hypothèques, lignes de crédit, autres billets (Digital Realty, Extra Space Storage)
    "dette_obligations": ["SeniorNotes"],
    "dette_non_garantie": ["UnsecuredDebt"],
    "dette_garantie": ["SecuredDebt"],
    "dette_ligne_credit": ["LineOfCredit", "LongTermLineOfCredit"],
    "dette_autres_billets": ["OtherNotesPayable"],
    "creances": ["AccountsReceivableNetCurrent", "ReceivablesNetCurrent", "AccountsNotesAndLoansReceivableNetCurrent"],
    "immobilisations": ["PropertyPlantAndEquipmentNet",
                        "PropertyPlantAndEquipmentAndFinanceLeaseRightOfUseAssetAfterAccumulatedDepreciationAndAmortization"],
}
INTERETS = ("charges_interets", "interets_nets", "interets_payes")
COMPOSANTES_DETTE = ("dette_obligations", "dette_non_garantie", "dette_garantie", "dette_ligne_credit",
                     "dette_autres_billets")
# Charges et décaissements, positifs par nature : certains dépôts en inversent le signe (Disney, Fortive)
CONCEPTS_POSITIFS = {"charges_interets", "interets_payes", "capex", "amortissements", "sga", "cout_revenus",
                     "couts_totaux"}
FLUX = {"revenus", "cout_revenus", "marge_brute", "couts_totaux", "resultat_exploitation", "resultat_avant_impot",
        "resultat_net", "impot", "charges_interets", "interets_nets", "interets_payes", "flux_dette",
        "cfo", "capex", "amortissements", "sga"}
STOCK = set(TAGS) - FLUX
FORMULAIRES = {"10-K", "10-K/A", "10-Q", "10-Q/A", "10-KT", "10-KT/A", "10-QT", "10-QT/A"}
COLONNES_FAITS = ["concept", "rang", "debut", "fin", "val", "depot", "form", "accn"]
# Sociétés passées sous une nouvelle holding (ou redomiciliées), avec un nouveau numéro CIK : leur historique
# comptable reste sous l'ancien numéro, dont elles prolongent les comptes. {nouveau CIK : ancien CIK}
PREDECESSEURS: Dict[int, int] = {
    2115436: 34088,      # ExxonMobil Holdings Corporation <- Exxon Mobil Corporation (juillet 2026)
    2012383: 1364742,    # BlackRock, Inc. (holding créée en 2024) <- BlackRock Finance, Inc. (ex-BlackRock, Inc.)
    1652044: 1288776,    # Alphabet (2015) <- Google
    1744489: 1001039,    # Walt Disney (2019) <- TWDC Enterprises 18 (ex-Walt Disney)
    1613103: 64670,      # Medtronic plc (2015) <- Medtronic, Inc.
    1730168: 1649338,    # Broadcom Inc. (2018) <- Broadcom Limited
    1649338: 1441634,    # Broadcom Limited (2016) <- Avago Technologies
    1707925: 884905,     # Linde plc (2018) <- Praxair
    1551182: 31277,      # Eaton Corporation plc (2012) <- Eaton Corporation
}
# Empreinte des règles d'extraction : les faits mis en cache avec d'autres étiquettes sont relus à la SEC
EMPREINTE_EXTRACTION = hashlib.sha256(json.dumps(
    {"tags": TAGS, "formulaires": sorted(FORMULAIRES), "colonnes": COLONNES_FAITS}, sort_keys=True
).encode()).hexdigest()[:10]


def extraire_faits(js: dict) -> pd.DataFrame:
    """Transforme le JSON « companyfacts » de la SEC en table compacte."""
    faits = js.get("facts", {}) if isinstance(js, dict) else {}
    usg = faits.get("us-gaap", {}) or {}
    dei = faits.get("dei", {}) or {}
    lignes = []

    def ajouter(concept: str, rang: int, noeud: dict, unite: str):
        for r in (noeud.get("units", {}).get(unite) or []):
            if r.get("form") not in FORMULAIRES:
                continue
            lignes.append((concept, rang, r.get("start"), r.get("end"), r.get("val"),
                           r.get("filed"), r.get("form"), r.get("accn")))

    for concept, tags in TAGS.items():
        for rang, tag in enumerate(tags):
            noeud = usg.get(tag)
            if noeud:
                ajouter(concept, rang, noeud, "USD")
    if dei.get("EntityCommonStockSharesOutstanding"):
        ajouter("actions", 0, dei["EntityCommonStockSharesOutstanding"], "shares")
    if usg.get("CommonStockSharesOutstanding"):           # total au bilan (sociétés multi-classes)
        ajouter("actions_bilan", 0, usg["CommonStockSharesOutstanding"], "shares")
    if usg.get("WeightedAverageNumberOfDilutedSharesOutstanding"):
        ajouter("actions_moy", 0, usg["WeightedAverageNumberOfDilutedSharesOutstanding"], "shares")
    df = pd.DataFrame(lignes, columns=COLONNES_FAITS)
    if df.empty:
        return df
    for c in ("debut", "fin", "depot"):
        df[c] = pd.to_datetime(df[c], errors="coerce")
    df["val"] = pd.to_numeric(df["val"], errors="coerce")
    df["rang"] = df["rang"].astype(int)
    return df.dropna(subset=["fin", "depot", "val"]).reset_index(drop=True)


class ClientSEC:
    """Accès poli à l'API de la SEC (10 requêtes/s max, adresse de contact obligatoire)."""
    BASE = "https://data.sec.gov"

    def __init__(self, cfg: Config):
        if not cfg.email_valide():
            raise ValueError(
                "La SEC exige une adresse e-mail de contact dans l'en-tête des requêtes. "
                "Renseigne ta propre adresse : Config(email='prenom.nom@exemple.fr').")
        self.cfg = cfg
        self.entetes = {"User-Agent": f"ScreenerSP500 {cfg.email.strip()}",
                        "Accept-Encoding": "gzip, deflate"}
        self._dernier = 0.0
        self._carte = None
        self._sic = None

    def _json(self, url: str) -> dict:
        attente = self.cfg.pause_sec - (time.time() - self._dernier)
        if attente > 0:
            time.sleep(attente)
        self._dernier = time.time()
        return http_get(url, self.entetes, timeout=60).json()

    def table_tickers(self) -> Dict[str, Tuple[int, str]]:
        if self._carte is None:
            chemin = self.cfg.chemin("cache", "sec_tickers.json")
            if _frais(chemin, 7):
                with open(chemin, encoding="utf-8") as fh:
                    data = json.load(fh)
            else:
                data = self._json("https://www.sec.gov/files/company_tickers.json")
                with open(chemin, "w", encoding="utf-8") as fh:
                    json.dump(data, fh)
            carte = {}
            for v in data.values():
                carte.setdefault(cle_ticker(v.get("ticker", "")), (int(v["cik_str"]), v.get("title", "")))
            self._carte = carte
        return self._carte

    def faits(self, cik: int) -> pd.DataFrame:
        chemin = self.cfg.chemin("cache", "sec", f"faits_{int(cik):010d}_{EMPREINTE_EXTRACTION}.pkl")
        if _frais(chemin, self.cfg.cache_sec_jours):
            return pd.read_pickle(chemin)
        anciens = sorted(glob.glob(self.cfg.chemin("cache", "sec", f"faits_{int(cik):010d}*.pkl")),
                         key=os.path.getmtime)
        try:
            js = self._json(f"{self.BASE}/api/xbrl/companyfacts/CIK{int(cik):010d}.json")
            df = extraire_faits(js)
        except RuntimeError as e:
            if "HTTP 404" not in str(e):
                if anciens:
                    info(f"  SEC indisponible pour le CIK {cik} : ancienne copie utilisée.")
                    return pd.read_pickle(anciens[-1])
                raise
            df = pd.DataFrame(columns=COLONNES_FAITS)
        df.to_pickle(chemin)
        for f in anciens:          # copies faites avec d'anciennes règles d'extraction
            if os.path.abspath(f) != os.path.abspath(chemin):
                os.remove(f)
        return df

    def sic(self, cik: int) -> Optional[int]:
        chemin = self.cfg.chemin("cache", "sec_sic.json")
        if self._sic is None:
            self._sic = {}
            if os.path.exists(chemin):
                with open(chemin, encoding="utf-8") as fh:
                    self._sic = json.load(fh)
        cle = str(int(cik))
        if cle not in self._sic:
            try:
                js = self._json(f"{self.BASE}/submissions/CIK{int(cik):010d}.json")
                self._sic[cle] = int(js.get("sic") or 0) or None
            except Exception:  # noqa: BLE001
                self._sic[cle] = None
            with open(chemin, "w", encoding="utf-8") as fh:
                json.dump(self._sic, fh)
        return self._sic[cle]


# ===========================================================================
# 5. Moteur « point-in-time » : un tableau par société, une ligne par trimestre
# ===========================================================================
# Chaque valeur est prise telle que publiée la première fois et n'est utilisée
# qu'à partir de sa date de dépôt. Les flux sont ramenés sur 12 mois glissants
# (TTM = dernier exercice + cumul de l'année en cours - cumul de l'an passé).

EPOQUE = pd.Timestamp("1970-01-01")


def _en_jours(serie: pd.Series) -> pd.Series:
    return (pd.to_datetime(serie, errors="coerce") - EPOQUE).dt.days


def _en_dates(jours) -> pd.DatetimeIndex:
    return pd.DatetimeIndex(pd.to_datetime(np.asarray(jours, dtype=float), unit="D", origin="unix")
                            ).astype("datetime64[ns]")


def _proches(index: dict, jour: int, tol: int) -> list:
    res = []
    get = index.get
    for j in range(jour - tol, jour + tol + 1):
        lst = get(j)
        if lst:
            res.extend(lst)
    return res


def _ttm(index: dict, jE: int, lim: int, revenus: bool = False, tol: int = 5, _niveau: int = 0) -> float:
    """Cumul sur 12 mois glissants de la période close vers jE, tel que connu à la date lim :
    exercice précédent + cumul de l'exercice en cours - même cumul un an plus tôt.

    La fin de période la plus proche et le cumul le plus long sont essayés d'abord ; s'ils ne permettent
    pas le calcul (trimestres « calendaires » ajoutés au rapport annuel, par exemple), les suivants
    prennent le relais. Si l'exercice précédent manque (date de clôture changée, période de
    transition), il est remplacé par le cumul sur 12 mois à sa date, calculé de la même façon."""
    # éléments : (jour_début, jour_fin, durée, valeur, jour_dépôt, rang_étiquette)
    cands = [p for p in _proches(index, jE, tol) if p[4] <= lim and 80 <= p[2] <= 380]
    if not cands:
        return np.nan
    pref = (lambda q: -q[3]) if revenus else (lambda q: q[5])
    for p in sorted(cands, key=lambda q: (abs(q[1] - jE), q[1], -q[2], pref(q))):
        if p[2] >= 350:
            return p[3]
        jd, fin, dur, val = p[0], p[1], p[2], p[3]
        fy = [q for q in _proches(index, jd - 1, 12) if 350 <= q[2] <= 380 and q[4] <= lim]
        if fy:
            a = min(fy, key=lambda q: (abs(q[1] - (jd - 1)), pref(q)))[3]
        elif _niveau == 0:
            a = _ttm(index, jd - 1, lim, revenus, 12, 1)
        else:
            continue
        cible = fin - 364
        ytd = [q for q in _proches(index, cible, 21) if abs(q[2] - dur) <= 20 and q[4] <= lim]
        if ytd and np.isfinite(a):
            b = min(ytd, key=lambda q: (abs(q[1] - cible), abs(q[2] - dur), pref(q)))
            return a + val - b[3]
    return np.nan


def _annuel_recent(index: dict, jE: int, lim: int, recul: int = 370) -> float:
    """Dernier exercice complet clos dans l'année précédant jE et publié à la date lim.

    Repli pour les montants que certaines sociétés ne publient que dans leur rapport annuel (charges
    d'intérêts d'Eli Lilly depuis 2024, par exemple) : faute de cumul sur 12 mois glissants, on reprend
    celui du dernier exercice, à condition qu'il ait moins d'un an."""
    if not index:
        return np.nan
    for j in range(jE, jE - recul - 1, -1):
        lst = index.get(j)
        if lst:
            ok = [q for q in lst if q[4] <= lim and 350 <= q[2] <= 380]
            if ok:
                return min(ok, key=lambda q: (q[5], q[4]))[3]
    return np.nan


def _premier_fini(*series: np.ndarray) -> np.ndarray:
    """Élément par élément, la première valeur finie de la liste (NaN si aucune)."""
    out = np.asarray(series[0], dtype=float)
    for x in series[1:]:
        out = np.where(np.isfinite(out), out, x)
    return out


def _trimestre(index: dict, jE: int, lim: int, tol: int = 5) -> float:
    """Valeur du seul trimestre clos à jE (T4 = exercice - 9 mois)."""
    cands = [p for p in _proches(index, jE, tol) if p[4] <= lim]
    if not cands:
        return np.nan
    fin = min(cands, key=lambda p: abs(p[1] - jE))[1]
    cands = [p for p in cands if p[1] == fin]
    q3 = [p for p in cands if 80 <= p[2] <= 100]
    if q3:
        return min(q3, key=lambda p: (p[5], p[4]))[3]
    longs = [p for p in cands if 150 <= p[2] <= 380]
    if not longs:
        return np.nan
    p = min(longs, key=lambda q: (-q[2], q[5]))
    prec = [x for x in _proches(index, fin - 91, 30)       # jusqu'à 121 jours : 4e trimestre de 16 semaines
            if x[4] <= lim and abs(x[0] - p[0]) <= 5 and 60 <= p[2] - x[2] <= 120]
    if not prec:
        return np.nan
    x = min(prec, key=lambda q: (abs(q[1] - (fin - 91)), q[5]))
    return p[3] - x[3]


def _instant(index: dict, jE: int, lim: int, tol: int = 7) -> float:
    # éléments : (valeur, jour_dépôt, rang_étiquette)
    for d in range(0, tol + 1):
        for j in ((jE,) if d == 0 else (jE - d, jE + d)):
            lst = index.get(j)
            if not lst:
                continue
            ok = [x for x in lst if x[1] <= lim]
            if ok:
                return min(ok, key=lambda x: (x[2], x[1]))[0]
    return np.nan


def _nb_actions(dei: list, bilan: dict, moy: dict, jE: int, lim: int, tol: int = 7) -> Tuple[float, float]:
    """Nombre d'actions et date de mesure : page de garde, sinon total du bilan, sinon moyenne diluée."""
    c = [x for x in dei if x[1] <= lim and jE - 10 <= x[0] <= jE + 200]
    if c:
        x = min(c, key=lambda x: x[1])
        return x[2], float(x[0])
    b = _instant(bilan, jE, lim, tol)
    if np.isfinite(b) and b > 0:
        return b, float(jE)
    m = [p for p in _proches(moy, jE, tol) if p[4] <= lim and 80 <= p[2] <= 380]
    if m:
        return min(m, key=lambda p: (abs(p[1] - jE), p[2], p[5]))[3], float(jE)
    return np.nan, np.nan


def _div(a, b):
    a = np.asarray(a, dtype=float)
    b = np.asarray(b, dtype=float)
    with np.errstate(divide="ignore", invalid="ignore"):
        out = np.where(np.isfinite(b) & (b != 0), a / b, np.nan)
    return out


def _div_pos(a, b):
    b = np.asarray(b, dtype=float)
    return np.where(b > 0, _div(a, b), np.nan)


def table_fondamentale(faits: pd.DataFrame) -> pd.DataFrame:
    """Une ligne par trimestre publié, datée par sa disponibilité ('dispo')."""
    if faits is None or len(faits) == 0:
        return pd.DataFrame()
    f = faits.copy()
    f["jd"] = _en_jours(f["debut"])
    f["jf"] = _en_jours(f["fin"])
    f["jp"] = _en_jours(f["depot"])
    f = f.dropna(subset=["jf", "jp", "val"])
    f["dur"] = f["jf"] - f["jd"]

    # Périodes publiées : la date de fin la plus récente de chaque dépôt 10-Q / 10-K
    duree = f[f["jd"].notna() & f["dur"].between(80, 400) & f["concept"].isin(FLUX)]
    if duree.empty:
        return pd.DataFrame()
    dep = duree.groupby("accn").agg(jf=("jf", "max"), jp=("jp", "min"), form=("form", "first"))
    dep["annuel"] = dep["form"].astype(str).str.startswith("10-K")
    per = dep.groupby("jf").agg(jp=("jp", "min"), annuel=("annuel", "max")).sort_index()
    per = per[(per["jp"].values - per.index.values) < 400]
    if per.empty:
        return pd.DataFrame()

    pos = f["concept"].isin(CONCEPTS_POSITIFS)
    f.loc[pos, "val"] = f.loc[pos, "val"].abs()

    # Première publication de chaque valeur, puis index de recherche par date de fin
    f = f.sort_values("jp", kind="mergesort")
    prem = f.drop_duplicates(subset=["concept", "rang", "jd", "jf"], keep="first")
    flux = {c: {} for c in FLUX}
    stock = {c: {} for c in STOCK}
    moy: dict = {}
    actions_bilan: dict = {}
    dei_accn: Dict[str, list] = {}
    for c, rang, jd, jf, val, jp, accn in zip(prem["concept"], prem["rang"], prem["jd"], prem["jf"],
                                              prem["val"], prem["jp"], prem["accn"]):
        jf_i, jp_i = int(jf), int(jp)
        if c in FLUX or c == "actions_moy":
            if jd != jd:  # NaN : valeur instantanée inattendue
                continue
            jd_i = int(jd)
            cible = moy if c == "actions_moy" else flux[c]
            cible.setdefault(jf_i, []).append((jd_i, jf_i, jf_i - jd_i, float(val), jp_i, int(rang)))
        elif c == "actions":
            e = dei_accn.setdefault(str(accn), [jf_i, jp_i, 0.0])
            e[0], e[1], e[2] = max(e[0], jf_i), min(e[1], jp_i), e[2] + float(val)
        elif c == "actions_bilan":
            actions_bilan.setdefault(jf_i, []).append((float(val), jp_i, 0))
        elif c in STOCK:
            stock[c].setdefault(jf_i, []).append((float(val), jp_i, int(rang)))
    dei = [tuple(v) for v in dei_accn.values()]
    fd = [p for lst in flux["flux_dette"].values() for p in lst]
    fd_fin = np.array([p[1] for p in fd], dtype=float)
    fd_dep = np.array([p[4] for p in fd], dtype=float)
    fd_pos = np.array([p[3] > 0 for p in fd], dtype=bool)

    def preuve_dette(jx: int, lim: int) -> float:
        """1 si un emprunt a été émis ou remboursé dans l'année close à jx (connu à la date lim), sinon 0."""
        if not len(fd_fin):
            return 0.0
        return float(np.any((fd_fin >= jx - 380) & (fd_fin <= jx + 5) & (fd_dep <= lim) & fd_pos))

    def valeurs(jx: int, lim: int, tol_flux: int, tol_bilan: int) -> dict:
        r = {c: _ttm(flux[c], jx, lim, revenus=(c == "revenus"), tol=tol_flux) for c in FLUX if c != "flux_dette"}
        for c in INTERETS:     # parfois publiés seulement dans le rapport annuel : repli sur le dernier exercice
            r[f"{c}_an"] = np.nan if np.isfinite(r[c]) else _annuel_recent(flux[c], jx, lim)
        r["flux_dette"] = preuve_dette(jx, lim)
        r.update({c: _instant(stock[c], jx, lim, tol_bilan) for c in STOCK})
        r["q_rn"] = _trimestre(flux["resultat_net"], jx, lim, tol_flux)
        r["actions"], r["j_actions"] = _nb_actions(dei, actions_bilan, moy, jx, lim, tol_bilan)
        return r

    # Une ligne par période publiée. Les valeurs « un an avant » viennent des comparatifs déjà publiés
    # à la date du dépôt (y compris ceux du dépôt lui-même) : utilisables dès le premier dépôt XBRL.
    lignes = []
    for jf, ligne in per.iterrows():
        jf_i, jp_i = int(jf), int(ligne["jp"])
        lim = jp_i + 30
        r = {"jf": jf_i, "jp": jp_i, "annuel": bool(ligne["annuel"])}
        r.update(valeurs(jf_i, lim, 5, 7))
        # un an et deux ans plus tôt : jusqu'à 35 jours d'écart, pour suivre un changement de date de clôture
        # (les fins de trimestre sont espacées d'environ 91 jours, aucune confusion possible)
        r.update({f"{c}__l1": x for c, x in valeurs(jf_i - 365, lim, 35, 35).items()})
        a2 = _instant(stock["actif"], jf_i - 730, lim, 35)
        r["actif__l2"] = a2 if np.isfinite(a2) else _instant(stock["passif_et_cp"], jf_i - 730, lim, 35)
        lignes.append(r)
    t = pd.DataFrame(lignes).sort_values("jf").reset_index(drop=True)
    n = len(t)
    v = {c: t[c].astype(float).values for c in t.columns if c not in ("annuel",)}

    def composer(g) -> Dict[str, np.ndarray]:
        rev = _premier_fini(g("revenus"), g("resultat_exploitation") + g("couts_totaux"))
        cp = g("capitaux_propres")
        # Charges d'intérêts : déclarées, sinon charge nette d'intérêts, sinon intérêts payés (12 mois
        # glissants), sinon les mêmes sur le dernier exercice publié s'il a moins d'un an ; à défaut, des
        # produits d'intérêts nets signifient qu'il n'y a pas de charge nette (0)
        def charge_nette(x):
            return np.where(np.isfinite(x) & (x < 0), -x, np.nan)

        def produit_net(x):
            return np.where(np.isfinite(x) & (x >= 0), 0.0, np.nan)
        interet = _premier_fini(g("charges_interets"), charge_nette(g("interets_nets")), g("interets_payes"),
                                g("charges_interets_an"), charge_nette(g("interets_nets_an")),
                                g("interets_payes_an"), produit_net(g("interets_nets")),
                                produit_net(g("interets_nets_an")))
        # EBIT : résultat d'exploitation, sinon avant impôt + intérêts, sinon résultat net + impôt + intérêts
        # (sans ligne d'impôt, comme les foncières qui n'en paient pas : impôt nul)
        ebit = np.where(np.isfinite(g("resultat_exploitation")), g("resultat_exploitation"),
                        np.where(np.isfinite(g("resultat_avant_impot")),
                                 g("resultat_avant_impot") + np.nan_to_num(interet),
                                 g("resultat_net") + np.nan_to_num(g("impot")) + np.nan_to_num(interet)))
        # Dette = part à plus d'un an + court terme (total déclaré, sinon échéances + emprunts),
        # sinon les agrégats. Rien au bilan : nulle, sauf indice contraire (intérêts, emprunts émis ou remboursés).
        nc, ech, ct_tot, ct = g("dette_lt_nc"), g("dette_lt_courante"), g("dette_courante"), g("emprunts_ct")
        court = np.where(np.isfinite(ct_tot), ct_tot, np.nan_to_num(ech) + np.nan_to_num(ct))
        lt_tot, tot = g("dette_lt_totale"), g("dette_totale")
        compo = [g(c) for c in COMPOSANTES_DETTE]
        somme_compo = np.where(np.any([np.isfinite(x) for x in compo], axis=0),
                               np.nansum([np.nan_to_num(x) for x in compo], axis=0), np.nan)
        dette = np.where(np.isfinite(nc), nc + court,
                         np.where(np.isfinite(lt_tot), lt_tot + np.nan_to_num(ct),
                                  np.where(np.isfinite(tot), tot,
                                           np.where(np.isfinite(somme_compo), somme_compo + np.nan_to_num(ct),
                                                    np.where(np.isfinite(ct_tot) | np.isfinite(ech)
                                                             | np.isfinite(ct), court, np.nan)))))
        # (des intérêts minimes, par exemple sur des contrats de location, ne prouvent pas une dette notable)
        seuil = np.where(np.isfinite(rev) & (rev > 0), 0.002 * rev, 0.001 * np.nan_to_num(g("actif")))
        preuve = (np.nan_to_num(interet) > seuil) | (np.nan_to_num(g("flux_dette")) > 0)
        dette = np.where(np.isfinite(dette), dette, np.where(preuve, np.nan, 0.0))
        return {
            "rev": rev,
            "gp": np.where(np.isfinite(g("marge_brute")), g("marge_brute"), rev - g("cout_revenus")),
            "ebit": ebit,
            "ni": g("resultat_net"), "cfo": g("cfo"), "capex": g("capex"), "da": g("amortissements"),
            "interet": interet, "sga": g("sga"), "actif": _premier_fini(g("actif"), g("passif_et_cp")),
            "ac": g("actif_courant"), "pc": g("passif_courant"), "cp": cp,
            "passif": np.where(np.isfinite(g("passif")), g("passif"), g("passif_et_cp") - cp),
            "treso": g("tresorerie") + np.nan_to_num(g("placements_ct")),
            "dette": dette,
            "creances": g("creances"), "immo": g("immobilisations"), "actions": g("actions"), "q_rn": g("q_rn"),
        }

    C = composer(lambda c: v[c])
    P = composer(lambda c: v[f"{c}__l1"])          # même trimestre, un an plus tôt
    rev, gp, ebit, ni, cfo, capex, da = C["rev"], C["gp"], C["ebit"], C["ni"], C["cfo"], C["capex"], C["da"]
    interet, sga, actif, ac, pc, cp = C["interet"], C["sga"], C["actif"], C["ac"], C["pc"], C["cp"]
    passif, treso, dette, creances, immo = C["passif"], C["treso"], C["dette"], C["creances"], C["immo"]
    actions, q_rn = C["actions"], C["q_rn"]
    jfs = t["jf"].values.astype(float)
    # Dette introuvable alors qu'elle existe (intérêts versés…) mais publiée au dernier rapport annuel
    # (Caterpillar, Lennar ne la détaillent pas chaque trimestre) : dernier montant connu de moins d'un an
    dette = dette.copy()
    derniere, j_derniere = np.nan, -1e9
    for k in range(n):
        if np.isfinite(dette[k]):
            derniere, j_derniere = dette[k], jfs[k]
        elif np.isfinite(derniere) and jfs[k] - j_derniere <= 370:
            dette[k] = derniere

    fcf = cfo - capex
    ebitda = ebit + np.nan_to_num(da)
    dette_nette = dette - treso
    actif_l1, actif_l2 = P["actif"], v["actif__l2"]
    actif_moy = np.where(np.isfinite(actif_l1), (actif + actif_l1) / 2, actif)

    # --- F-score de Piotroski (9 signaux, calculés en glissement annuel) ---
    roa = _div_pos(ni, actif_l1)
    roa_l1 = _div_pos(P["ni"], actif_l2)
    cfo_ta = _div_pos(cfo, actif_l1)
    levier = _div_pos(dette, actif)
    levier_l1 = _div_pos(P["dette"], actif_l1)
    cr, cr_l1 = _div_pos(ac, pc), _div_pos(P["ac"], P["pc"])
    gm, gm_l1 = _div_pos(gp, rev), _div_pos(P["gp"], P["rev"])
    at = _div_pos(rev, actif_l1)
    at_l1 = _div_pos(P["rev"], actif_l2)
    act_l1 = P["actions"]
    ratio_act = _div(actions, act_l1)

    def sig(cond, *entrees):
        ok = np.all([np.isfinite(e) for e in entrees], axis=0)
        return np.where(ok, cond.astype(float), np.nan)

    signaux = np.vstack([
        sig(roa > 0, roa),
        sig(cfo > 0, cfo),
        sig(roa > roa_l1, roa, roa_l1),
        sig(cfo_ta > roa, cfo_ta, roa),
        sig((levier < levier_l1) | ((levier == 0) & (levier_l1 == 0)), levier, levier_l1),
        sig(cr > cr_l1, cr, cr_l1),
        # pas d'émission d'actions ; un saut > 40 % ou une baisse > 30 % = division d'actions -> inconnu
        sig(actions <= act_l1 * 1.005, np.where((ratio_act > 1.4) | (ratio_act < 0.7), np.nan, actions), act_l1),
        sig(gm > gm_l1, gm, gm_l1),
        sig(at > at_l1, at, at_l1),
    ])
    fscore_n = np.isfinite(signaux).sum(axis=0).astype(float)
    fscore = np.nansum(signaux, axis=0)
    fscore = np.where(fscore_n > 0, fscore, np.nan)

    # --- M-score de Beneish (8 variables) ---
    with np.errstate(divide="ignore", invalid="ignore"):
        dsri = _div(_div_pos(creances, rev), _div_pos(P["creances"], P["rev"]))
        gmi = _div(gm_l1, gm)
        aqi = _div(1 - _div_pos(ac + immo, actif), 1 - _div_pos(P["ac"] + P["immo"], actif_l1))
        sgi = _div_pos(rev, P["rev"])
        depi = _div(_div_pos(P["da"], P["da"] + P["immo"]), _div_pos(da, da + immo))
        sgai = _div(_div_pos(sga, rev), _div_pos(P["sga"], P["rev"]))
        lvgi = _div(_div_pos(pc + dette, actif), _div_pos(P["pc"] + P["dette"], actif_l1))
        tata = _div_pos(ni - cfo, actif)
        beneish = (-4.84 + 0.920 * dsri + 0.528 * gmi + 0.404 * aqi + 0.892 * sgi + 0.115 * depi
                   - 0.172 * sgai + 4.679 * tata - 0.327 * lvgi)

    # --- Historique annuel : stabilité de la rentabilité et FCF sur 3 ans ---
    annuel = t["annuel"].values.astype(bool)
    pos_ann = np.where(annuel)[0]
    stabilite = np.full(n, np.nan)
    fcf_3a = np.full(n, np.nan)
    for k in range(n):
        prec = pos_ann[pos_ann <= k]
        if len(prec) >= 3:
            vals = roa[prec[-5:]]
            vals = vals[np.isfinite(vals)]
            if len(vals) >= 3:
                stabilite[k] = -float(np.std(vals, ddof=1))
            f3 = fcf[prec[-3:]]
            if np.isfinite(f3).all() and (jfs[prec[-1]] - jfs[prec[-3]]) < 800:
                fcf_3a[k] = float(f3.sum())

    with np.errstate(divide="ignore", invalid="ignore"):
        dette_ebitda = np.where(ebitda > 0, dette_nette / ebitda,
                                np.where(np.isfinite(ebitda) & (dette_nette > 0), np.inf, np.nan))
        couverture = np.where(interet > 0, ebit / interet, np.nan)

    out = pd.DataFrame({
        "fin": _en_dates(jfs),
        "dispo": _en_dates(t["jp"].values),
        "annuel": annuel,
        "revenus": rev, "marge_brute": gp, "ebit": ebit, "resultat_net": ni, "cfo": cfo, "capex": capex,
        "fcf": fcf, "ebitda": ebitda, "charges_interets": interet, "actif": actif, "capitaux_propres": cp,
        "passif": passif, "dette": dette, "tresorerie": treso, "dette_nette": dette_nette, "actions": actions,
        "actions_date": _en_dates(v["j_actions"]),
        "fscore": fscore, "fscore_n": fscore_n, "couverture": couverture, "dette_ebitda": dette_ebitda,
        "accruals": _div_pos(ni - cfo, actif_moy), "beneish": beneish,
        "rentab_brute": _div_pos(gp, actif), "roa_moy": _div_pos(ni, actif_moy),
        "cfo_actif": _div_pos(cfo, actif_moy), "stabilite": stabilite, "fcf_3a": fcf_3a,
        "levier_fin": _div_pos(passif, cp), "surprise_rn": _div_pos(q_rn - P["q_rn"], actif),
        "marge_brute_pct": gm,
    })
    return out


COLONNES_DATES = ("fin", "dispo", "actions_date")
COLONNES_FOND = [c for c in [
    "fin", "dispo", "revenus", "marge_brute", "ebit", "resultat_net", "cfo", "capex", "fcf", "ebitda",
    "charges_interets", "actif", "capitaux_propres", "passif", "dette", "tresorerie", "dette_nette", "actions",
    "actions_date",
    "fscore", "fscore_n", "couverture", "dette_ebitda", "accruals", "beneish", "rentab_brute", "roa_moy",
    "cfo_actif", "stabilite", "fcf_3a", "levier_fin", "surprise_rn", "marge_brute_pct"]]


def sans_lignes_degradees(t: pd.DataFrame, recul_jours: int = 400) -> pd.DataFrame:
    """Retire les lignes incomplètes (F-score sur moins de 6 critères ou actif inconnu) publiées après une
    ligne complète portant sur une période close moins de 400 jours plus tôt : cette dernière reste en
    vigueur. Cas d'une scission récente (comparatifs trimestriels absents), d'un changement de date de
    clôture ou d'une étiquette abandonnée d'un rapport à l'autre."""
    if t is None or t.empty or "fscore_n" not in t:
        return t
    complet = ((t["fscore_n"] >= 6) & t["actif"].notna()).values
    if complet.all() or not complet.any():
        return t
    garder = np.ones(len(t), dtype=bool)
    fins = pd.to_datetime(t["fin"]).values
    dernier = None
    for k in np.argsort(pd.to_datetime(t["dispo"]).values, kind="stable"):
        if complet[k]:
            dernier = fins[k] if dernier is None else max(dernier, fins[k])
        elif dernier is not None and (fins[k] - dernier) / np.timedelta64(1, "D") <= recul_jours:
            garder[k] = False
    return t[garder]


def panneau_fondamental(tables: Dict[int, pd.DataFrame]) -> pd.DataFrame:
    morceaux = []
    for cik, t in tables.items():
        t = sans_lignes_degradees(t)
        if t is None or t.empty:
            continue
        x = t.copy()
        x["cik"] = int(cik)
        morceaux.append(x)
    if not morceaux:
        return pd.DataFrame()
    p = pd.concat(morceaux, ignore_index=True)
    p["dispo_eff"] = (p["dispo"] + pd.Timedelta(days=1)).astype("datetime64[ns]")
    p["cik"] = p["cik"].astype("int64")
    return p.sort_values("dispo_eff").reset_index(drop=True)


def fondamentaux_aux_dates(panneau: pd.DataFrame, dates: Sequence, ciks: Sequence[int],
                           fraicheur_jours: int = 450) -> pd.DataFrame:
    """Pour chaque date et chaque société : la dernière ligne disponible à cette date."""
    if panneau is None or panneau.empty or not len(ciks):
        return pd.DataFrame()
    dates = pd.DatetimeIndex(pd.to_datetime(list(dates))).astype("datetime64[ns]")
    ciks = np.array(sorted({int(c) for c in ciks}), dtype="int64")
    gauche = pd.DataFrame({"date": np.repeat(dates.values, len(ciks)), "cik": np.tile(ciks, len(dates))})
    gauche = gauche.sort_values("date", kind="mergesort").reset_index(drop=True)
    m = pd.merge_asof(gauche, panneau, left_on="date", right_on="dispo_eff", by="cik", direction="backward")
    vieux = ((m["date"] - m["fin"]).dt.days > fraicheur_jours) | m["fin"].isna()
    num = [c for c in m.columns if c not in ("date", "cik", "dispo_eff", "annuel") + COLONNES_DATES
           and pd.api.types.is_numeric_dtype(m[c])]
    m.loc[vieux, num] = np.nan
    for c in COLONNES_DATES:
        if c in m:
            m.loc[vieux, c] = pd.NaT
    return m


# ===========================================================================
# 6. Facteurs de prix
# ===========================================================================

def matrices_prix(close: pd.DataFrame, spy: pd.Series) -> Dict[str, pd.DataFrame]:
    ret = close.pct_change(fill_method=None)
    mom = close.shift(21) / close.shift(252) - 1.0
    vol = ret.rolling(252, min_periods=200).std() * np.sqrt(252)
    s = spy.reindex(close.index).pct_change(fill_method=None)
    masque = ret.notna().values & s.notna().values[:, None]
    X = ret.where(masque)
    Y = pd.DataFrame(np.where(masque, s.values[:, None], np.nan), index=close.index, columns=close.columns)
    roll = lambda d: d.rolling(252, min_periods=200).mean()   # noqa: E731
    mx, my = roll(X), roll(Y)
    cov = roll(X * Y) - mx * my
    var = roll(Y * Y) - my * my
    beta = cov / var.where(var > 0)
    return {"ret": ret, "mom": mom, "vol": vol, "beta": beta}


def perte_max_1an(close: pd.DataFrame, positions: Sequence[int]) -> pd.DataFrame:
    vals = close.values
    lignes = []
    with np.errstate(invalid="ignore", divide="ignore"):
        for i in positions:
            w = vals[max(0, i - 251): i + 1]
            plus_haut = np.fmax.accumulate(w, axis=0)
            dd = w / plus_haut - 1.0
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                lignes.append(np.nanmin(np.where(np.isfinite(dd), dd, np.nan), axis=0))
    return pd.DataFrame(lignes, index=close.index[list(positions)], columns=close.columns)


# ===========================================================================
# 7. Tableau des facteurs, filtres de sécurité, scores et signaux
# ===========================================================================

def _facteur_a_dates(facteur: Optional[pd.DataFrame], dates, tick: Sequence[str]) -> np.ndarray:
    out = np.ones(len(tick))
    if facteur is None or facteur.empty:
        return out
    idx = facteur.index
    vals = facteur.reindex(columns=list(tick)).values
    dd = pd.to_datetime(pd.Series(list(dates))).values
    for k in range(len(tick)):
        a = dd[k]
        if pd.isna(a):
            continue
        p = min(max(int(idx.searchsorted(a, side="right")) - 1, 0), len(idx) - 1)
        v = vals[p, k]
        out[k] = v if np.isfinite(v) and v > 0 else 1.0
    return out


def tableau_facteurs(d, i_d: int, tickers: Sequence[str], close: pd.DataFrame, mats: Dict[str, pd.DataFrame],
                     mdd_ligne: pd.Series, infos: pd.DataFrame, fond_d: Optional[pd.DataFrame],
                     revisions: Optional[pd.Series] = None, mode: str = "complet",
                     capi_externe: Optional[pd.Series] = None, close_split: Optional[pd.DataFrame] = None,
                     facteur: Optional[pd.DataFrame] = None) -> pd.DataFrame:
    tick = [t for t in tickers if t in close.columns]
    df = pd.DataFrame(index=pd.Index(tick, name="ticker"))
    df["prix"] = close.iloc[i_d][tick].values.astype(float)
    df["prix_split"] = (close_split.iloc[i_d].reindex(tick).values.astype(float)
                        if close_split is not None else df["prix"].values)
    nb = close.iloc[max(0, i_d - 251): i_d + 1][tick].notna().sum().values
    df["historique_ok"] = (nb >= 200) & np.isfinite(df["prix"].values)
    for c in ("mom", "vol", "beta"):
        df[c] = mats[c].iloc[i_d][tick].values.astype(float)
    df["mdd"] = mdd_ligne.reindex(tick).values.astype(float)
    inf_t = infos.reindex(tick)
    df["nom"] = inf_t["nom"].fillna("").astype(str).values
    df["secteur"] = inf_t["secteur"].fillna("Inconnu").astype(str).values
    df["cik"] = pd.to_numeric(inf_t["cik"], errors="coerce").values

    if mode == "complet":
        if fond_d is not None and len(fond_d):
            f = fond_d.drop_duplicates("cik", keep="last").set_index("cik")
            cles = [int(c) if np.isfinite(c) else -1 for c in df["cik"].values]
            jf = f.reindex(cles)
            for c in COLONNES_FOND:
                if c in jf:
                    df[c] = jf[c].values
        else:
            for c in COLONNES_FOND:
                df[c] = np.nan
        df["f_actions"] = _facteur_a_dates(facteur, df["actions_date"].values, tick)
        metriques_marche(df, d, capi_externe)
    if revisions is not None:
        df["revision"] = pd.to_numeric(revisions.reindex(tick), errors="coerce").values
    return df


def metriques_marche(df: pd.DataFrame, d, capi_externe: Optional[pd.Series] = None) -> None:
    """Capitalisation, valeur d'entreprise et rendements (FCF, EBIT, bénéfices), avec contrôles.
    Capitalisation = cours ajusté des divisions × actions publiées × divisions survenues depuis."""
    f_act = df["f_actions"].astype(float).values if "f_actions" in df else 1.0
    capi = df["prix_split"].astype(float).values * df["actions"].astype(float).values * f_act
    rn = df["resultat_net"].astype(float).values
    cp = df["capitaux_propres"].astype(float).values
    with np.errstate(divide="ignore", invalid="ignore"):
        suspect = (~np.isfinite(capi)) | (capi < 5e8) \
            | (np.abs(rn / capi) > 1.0) | (cp / capi > 25) \
            | np.isin(df.index.values, list(CAPI_NON_FIABLE_SEC))
    capi = np.where(suspect, np.nan, capi)
    if capi_externe is not None:
        ext = pd.to_numeric(capi_externe.reindex(df.index), errors="coerce").values
        capi = np.where(np.isfinite(ext) & (ext > 0), ext, capi)
    df["capitalisation"] = capi
    # dette inconnue -> valeur d'entreprise inconnue (plutôt que de la supposer nulle)
    ve = capi + df["dette"].astype(float).values - np.nan_to_num(df["tresorerie"].astype(float).values)
    df["rend_fcf"] = _div_pos(df["fcf"].astype(float).values, capi)
    df["ebit_ve"] = _div_pos(df["ebit"].astype(float).values, ve)
    df["rend_benef"] = _div_pos(rn, capi)
    df["bm"] = _div_pos(cp, capi)
    recent = (pd.Timestamp(d) - pd.to_datetime(df["fin"])).dt.days.values <= 200
    df["surprise"] = np.where(recent, df["surprise_rn"].astype(float).values, np.nan)


def rang_secteur(valeurs: pd.Series, secteurs: pd.Series, min_groupe: int = 5) -> pd.Series:
    """Percentile (0-1) de chaque titre au sein de son secteur (ou de l'indice si secteur trop petit)."""
    v = pd.to_numeric(valeurs, errors="coerce")
    glob = v.rank(pct=True)
    g = v.groupby(secteurs)
    rk = g.rank(pct=True)
    nb = g.transform("count")
    out = rk.where(nb >= min_groupe, glob)
    return out.where(v.notna())


def _moyenne(series: List[pd.Series], min_n: int) -> pd.Series:
    m = pd.concat(series, axis=1)
    return m.mean(axis=1).where(m.notna().sum(axis=1) >= min_n)


def _fr(x: float, dec: int = 1) -> str:
    return f"{x:.{dec}f}".replace(".", ",")


def appliquer_filtres(df: pd.DataFrame, cfg: Config) -> pd.DataFrame:
    df = df.copy()
    n = len(df)
    motifs: List[List[str]] = [[] for _ in range(n)]
    alertes: List[List[str]] = [[] for _ in range(n)]
    echec = np.zeros(n, dtype=bool)
    sect = df["secteur"].fillna("Inconnu")
    nonfin = (sect != "Financials").values

    def marquer(masque, texte, liste=motifs, bloque=True, cle=None, mesure=None, applicable=True):
        m = np.asarray(pd.Series(masque, index=df.index).fillna(False), dtype=bool)
        for k in np.where(m)[0]:
            liste[k].append(texte(k))
        if bloque:
            echec[m] = True
        if cle:   # résultat de chaque filtre, pour l'affichage : « echec », « ok » ou « na » (non mesuré)
            mes = np.isfinite(np.asarray(pd.to_numeric(pd.Series(mesure, index=df.index), errors="coerce"),
                                         dtype=float))
            app = np.asarray(pd.Series(applicable, index=df.index).fillna(False), dtype=bool)
            df[f"f_{cle}"] = np.where(m, "echec", np.where(app & mes, "ok", "na"))

    vol = df["vol"].astype(float)
    seuil_vol = vol.quantile(1 - cfg.exclusion_volatilite)
    pv, pl, pa = (f"{cfg.exclusion_volatilite * 100:g}", f"{cfg.exclusion_levier_secteur * 100:g}",
                  f"{cfg.exclusion_accruals_secteur * 100:g}")
    marquer(vol > seuil_vol, lambda k: f"Volatilité {_fr(vol.iloc[k] * 100, 0)} % "
                                       f"({pv} % les plus volatils de l'indice)", cle="volatilite", mesure=vol)
    df["s_volatilite"] = seuil_vol
    donnees_ok = df["historique_ok"].values.astype(bool)
    df["fscore_9"] = np.nan

    if cfg.mode == "complet":
        fn = df["fscore_n"].astype(float)
        fs9 = (df["fscore"].astype(float) * 9 / fn).where(fn >= criteres_fscore_requis(sect))
        df["fscore_9"] = fs9
        donnees_ok = donnees_ok & fs9.notna().values & df["actif"].astype(float).notna().values
        marquer(fs9 < cfg.fscore_min, lambda k: f"F-score {_fr(fs9.iloc[k])}/9 < {cfg.fscore_min:g}",
                cle="fscore", mesure=fs9)

        seuil_cov = sect.map(lambda s: cfg.couverture_min_secteurs.get(s, cfg.couverture_min)).astype(float)
        cov = df["couverture"].astype(float)
        marquer(nonfin & (cov < seuil_cov),
                lambda k: f"Couverture des intérêts {_fr(cov.iloc[k])}x < {_fr(seuil_cov.iloc[k])}x",
                cle="couverture", mesure=cov, applicable=nonfin)
        df["s_couverture"] = seuil_cov

        ebitda = df["ebitda"].astype(float)
        dn = df["dette_nette"].astype(float)
        marquer(nonfin & (ebitda <= 0) & (dn > 0), lambda k: "EBITDA négatif avec une dette nette",
                cle="ebitda", mesure=ebitda + dn, applicable=nonfin)
        de = df["dette_ebitda"].astype(float)
        de_fini = de.where(np.isfinite(de))
        q_de = de_fini.groupby(sect).transform(lambda x: x.quantile(1 - cfg.exclusion_levier_secteur)
                                               if x.notna().sum() >= 5 else np.nan)
        marquer(nonfin & np.isfinite(de) & (de > q_de),
                lambda k: f"Dette nette/EBITDA {_fr(de.iloc[k])} ({pl} % les plus endettés du secteur)",
                cle="levier", mesure=de_fini.where(q_de.notna()), applicable=nonfin)
        df["s_levier"] = q_de

        exempt = sect.isin(list(cfg.secteurs_sans_filtre_fcf)).values
        f3 = df["fcf_3a"].astype(float)
        f1 = df["fcf"].astype(float)
        marquer(nonfin & ~exempt & ((f3 <= 0) | (f3.isna() & (f1 <= 0))),
                lambda k: "Free cash-flow négatif (cumul 3 ans ou 12 derniers mois)",
                cle="fcf", mesure=f3.where(f3.notna(), f1), applicable=nonfin & ~exempt)

        # Accruals de Sloan : part du bénéfice non convertie en trésorerie. Sont exclus les 10 % les pires
        # du secteur, mais seulement si le bénéfice dépasse le cash-flow d'exploitation (accruals > 0) :
        # dans les secteurs lourds en amortissements, le décile « pire » peut rester négatif, et une
        # société dont le cash-flow couvre déjà le bénéfice n'a pas de bénéfice gonflé à craindre.
        acc = df["accruals"].astype(float)
        q_acc = acc.groupby(sect).transform(lambda x: x.quantile(1 - cfg.exclusion_accruals_secteur)
                                            if x.notna().sum() >= 5 else np.nan)
        seuil_acc = q_acc.clip(lower=0.0)
        marquer(nonfin & (acc > seuil_acc),
                lambda k: f"Accruals élevés ({_fr(acc.iloc[k] * 100)} % de l'actif, {pa} % les pires du secteur)",
                cle="accruals", mesure=acc.where(q_acc.notna()), applicable=nonfin)
        df["s_accruals"] = seuil_acc

        ben = df["beneish"].astype(float)
        marquer(nonfin & (ben > -1.78),
                lambda k: f"M-score de Beneish {_fr(ben.iloc[k], 2)} > -1,78",
                liste=motifs if cfg.beneish_exclut else alertes, bloque=cfg.beneish_exclut,
                cle="beneish", mesure=ben, applicable=nonfin)

    df["echec_filtre"] = echec
    df["donnees_ok"] = donnees_ok
    df["motif"] = ["; ".join(m) for m in motifs]
    df["alertes"] = ["; ".join(a) for a in alertes]
    return df


def calculer_scores(df: pd.DataFrame, cfg: Config) -> pd.DataFrame:
    df = df.copy()
    sect = df["secteur"].fillna("Inconnu")
    est_fin = sect == "Financials"
    non_fin = ~est_fin
    nan = pd.Series(np.nan, index=df.index)

    def rg(col, sens=1.0, masque=None):
        if col not in df:
            return nan
        v = pd.to_numeric(df[col], errors="coerce") * sens
        if masque is not None:
            v = v.where(masque)
        return rang_secteur(v, sect)

    p = {"momentum": rg("mom")}
    res = [rg("surprise")]
    if "revision" in df:
        res.append(rg("revision"))
    p["resultats"] = _moyenne(res, 1)
    if cfg.mode == "complet":
        p["valeur"] = _moyenne([rg("rend_fcf", masque=non_fin), rg("ebit_ve", masque=non_fin),
                                rg("rend_benef"), rg("bm", masque=est_fin)], 1)
        levier = rg("dette_ebitda", -1.0, non_fin).combine_first(rg("levier_fin", -1.0, est_fin))
        p["qualite"] = _moyenne([rg("rentab_brute", masque=non_fin), rg("roa_moy"),
                                 rg("cfo_actif", masque=non_fin), rg("stabilite"), levier], 2)
    else:
        p["valeur"], p["qualite"] = nan, nan
    p["faible_risque"] = _moyenne([rg("vol", -1.0), rg("beta", -1.0), rg("mdd")], 2)

    min_pot = 2 if cfg.mode == "complet" else 1
    pot = _moyenne([p["momentum"], p["resultats"], p["valeur"]], min_pot).where(p["momentum"].notna())
    if cfg.mode == "complet":
        rob = _moyenne([p["qualite"], p["faible_risque"]], 2)
    else:
        rob = p["faible_risque"]
    glob = cfg.poids_potentiel * pot + (1 - cfg.poids_potentiel) * rob
    df["potentiel_pct"] = pot.rank(pct=True) * 100
    df["robustesse_pct"] = rob.rank(pct=True) * 100
    df["score"] = glob.rank(pct=True) * 100
    for k, s in p.items():
        df[f"p_{k}"] = s * 100
    return df


SIGNAUX = ["ACHETER", "GARDER", "SURVEILLER", "VENDRE", "DONNÉES INSUFFISANTES"]


def criteres_fscore_requis(secteur):
    """Critères du F-score à mesurer au minimum : les deux tiers des critères applicables, soit 6 sur 9,
    et 5 sur 7 pour les financières (ni ratio de liquidité ni marge brute dans leurs comptes)."""
    if isinstance(secteur, (pd.Series, np.ndarray, list)):
        return np.where(np.asarray(secteur) == "Financials", 5, 6)
    return 5 if secteur == "Financials" else 6


def _raison_insuffisance(r: pd.Series, cfg: Config) -> str:
    """Pourquoi une action n'a pas de signal : le premier manque rencontré."""
    if not bool(r.get("historique_ok", False)):
        return "Moins de 200 séances de cours sur la dernière année (introduction en bourse ou scission récente)"
    if cfg.mode == "complet":
        fin = r.get("fin")
        if fin is None or pd.isna(fin):
            return "Aucun compte récent exploitable à la SEC (société récente, ou dernier rapport trop ancien)"
        n = pd.to_numeric(r.get("fscore_n"), errors="coerce")
        n = int(n) if np.isfinite(n) else 0
        requis = criteres_fscore_requis(r.get("secteur"))
        if n < requis:
            if n == 0:
                return "F-score non mesurable : comptes XBRL incomplets"
            return (f"F-score mesurable sur {n} critère{'s' if n > 1 else ''} seulement ({requis} requis) : "
                    "comptes XBRL incomplets")
        if not np.isfinite(pd.to_numeric(r.get("actif"), errors="coerce")):
            return "Total de l'actif introuvable dans les comptes XBRL"
    if not np.isfinite(pd.to_numeric(r.get("mom"), errors="coerce")):
        return "Moins d'un an de cours : momentum sur 12 mois incalculable (introduction en bourse ou scission récente)"
    return "Score incalculable : trop d'indicateurs manquants"


def attribuer_signaux(df: pd.DataFrame, cfg: Config) -> pd.DataFrame:
    df = df.copy()
    s = df["score"].astype(float)
    rob = df["robustesse_pct"].astype(float)
    donnees = df["donnees_ok"].values & s.notna().values
    conds = [~donnees, df["echec_filtre"].values,
             ((s >= cfg.seuil_achat) & (rob >= cfg.robustesse_min_achat)).values,
             (s >= cfg.seuil_garder).values, (s >= cfg.seuil_vente).values]
    df["signal"] = np.select(conds, ["DONNÉES INSUFFISANTES", "VENDRE", "ACHETER", "GARDER", "SURVEILLER"],
                             default="VENDRE")
    motif = df["motif"].astype(str).values.copy()
    for k in range(len(df)):
        sig = df["signal"].iloc[k]
        if sig == "VENDRE" and not motif[k]:
            motif[k] = f"Score {_fr(s.iloc[k], 0)} < {cfg.seuil_vente:g} (moitié basse du classement)"
        elif sig == "ACHETER":
            motif[k] = f"Top {100 - cfg.seuil_achat:g} % du classement, filtres de sécurité passés"
        elif sig == "DONNÉES INSUFFISANTES":
            raison = _raison_insuffisance(df.iloc[k], cfg)
            motif[k] = raison + (f"; {motif[k]}" if motif[k] else "")
    df["motif"] = motif
    return df


def construire_portefeuille(df: pd.DataFrame, actuels: Sequence[str], cfg: Config,
                            hors_univers: str = "Sorti de l'indice ou sans données"
                            ) -> Tuple[List[str], List[str], List[Tuple[str, str]]]:
    """Conserve les positions encore valables puis complète avec les meilleurs ACHETER."""
    n = cfg.n_positions
    plafond = max(1, int(math.floor(cfg.max_par_secteur * n + 1e-9)))
    gardes: List[str] = []
    ventes: List[Tuple[str, str]] = []
    for t in dict.fromkeys(actuels):
        if t not in df.index:
            ventes.append((t, hors_univers))
        elif df.at[t, "signal"] in ("ACHETER", "GARDER", "SURVEILLER", "DONNÉES INSUFFISANTES"):
            gardes.append(t)
        else:
            ventes.append((t, str(df.at[t, "motif"])))
    compte = Counter(df.loc[gardes, "secteur"]) if gardes else Counter()
    ciks = {c for c in df.loc[gardes, "cik"] if np.isfinite(c)} if gardes else set()
    achats: List[str] = []
    cands = df[(df["signal"] == "ACHETER") & ~df.index.isin(gardes)].sort_values("score", ascending=False)
    for t, ligne in cands.iterrows():
        if len(gardes) + len(achats) >= n:
            break
        if compte[ligne["secteur"]] >= plafond:
            continue
        if np.isfinite(ligne["cik"]) and ligne["cik"] in ciks:
            continue   # une seule classe d'actions par société (GOOGL / GOOG…)
        achats.append(t)
        compte[ligne["secteur"]] += 1
        if np.isfinite(ligne["cik"]):
            ciks.add(ligne["cik"])
    return gardes, achats, ventes


# ===========================================================================
# 8. Données Yahoo du jour : révisions d'analystes et capitalisations
# ===========================================================================

def _revision(et) -> float:
    if et is None or len(et) == 0:
        return np.nan
    et = et.copy()
    et.columns = [str(c).lower() for c in et.columns]
    et.index = [str(i).lower() for i in et.index]
    vals = []
    for per in ("0y", "+1y"):
        if per in et.index and "current" in et.columns and "90daysago" in et.columns:
            cur, old = et.at[per, "current"], et.at[per, "90daysago"]
            try:
                cur, old = float(cur), float(old)
            except (TypeError, ValueError):
                continue
            if np.isfinite(cur) and np.isfinite(old) and abs(old) > 1e-6:
                vals.append(float(np.clip((cur - old) / abs(old), -1, 1)))
    return float(np.mean(vals)) if vals else np.nan


def lire_revisions(tickers: Sequence[str], cfg: Config) -> pd.Series:
    """Variation sur 90 jours du BPA moyen attendu par les analystes (exercice en cours et suivant)."""
    import yfinance as yf

    chemin = cfg.chemin("cache", f"revisions_{aujourdhui():%Y%m%d}.pkl")
    vals = pd.read_pickle(chemin).to_dict() if os.path.exists(chemin) else {}
    restants = [t for t in tickers if t not in vals]
    if restants:
        info(f"Révisions d'analystes : {len(restants)} titres à interroger (2 à 5 minutes)…")
    echecs = 0
    for k, t in enumerate(restants):
        try:
            vals[t] = _revision(yf.Ticker(t).eps_trend)
            echecs = 0
        except Exception as e:  # noqa: BLE001
            vals[t] = np.nan
            echecs += 1
            texte = f"{type(e).__name__} {e}".lower()
            if "rate" in texte or "too many" in texte or echecs >= 15:
                info("  Yahoo limite les requêtes : révisions arrêtées ici (le classement les ignore).")
                break
        if (k + 1) % 100 == 0:
            info(f"  révisions : {k + 1}/{len(restants)}")
            pd.Series(vals, dtype=float).to_pickle(chemin)
        time.sleep(0.2)
    s = pd.Series(vals, dtype=float)
    s.to_pickle(chemin)
    ETAT_SOURCES["revisions"] = {"demandes": len(tickers), "obtenues": int(s.reindex(list(tickers)).notna().sum())}
    return s


def lire_capitalisations(tickers: Sequence[str]) -> pd.Series:
    """Capitalisation du jour chez Yahoo, pour les rares titres où la SEC ne suffit pas."""
    import yfinance as yf

    out = {}
    for t in tickers:
        try:
            v = yf.Ticker(t).info.get("marketCap")
            out[t] = float(v) if v else np.nan
        except Exception:  # noqa: BLE001
            out[t] = np.nan
        time.sleep(0.3)
    return pd.Series(out, dtype=float)


# ===========================================================================
# 9. Chargement des comptes pour une liste de sociétés
# ===========================================================================

def infos_tickers(tickers: Sequence[str], actuel: pd.DataFrame, sec: Optional[ClientSEC]) -> pd.DataFrame:
    base = actuel.set_index("ticker")[["nom", "secteur", "cik"]]
    carte = sec.table_tickers() if sec is not None else {}
    lignes = {}
    inconnus = [t for t in tickers if t not in base.index]
    for t in tickers:
        if t in base.index:
            r = base.loc[t]
            lignes[t] = (r["nom"], r["secteur"], r["cik"])
    if inconnus:
        info(f"Secteurs des anciens membres : {len(inconnus)} sociétés…")
    for t in inconnus:
        cik, titre = carte.get(cle_ticker(t), (np.nan, t))
        secteur = "Inconnu"
        if sec is not None and np.isfinite(cik):
            secteur = secteur_depuis_sic(sec.sic(int(cik)))
        lignes[t] = (titre, secteur, cik)
    out = pd.DataFrame.from_dict(lignes, orient="index", columns=["nom", "secteur", "cik"])
    out.index.name = "ticker"
    out["cik"] = pd.to_numeric(out["cik"], errors="coerce")
    return out


def fusionner_predecesseur(nouveau: pd.DataFrame, ancien: pd.DataFrame) -> pd.DataFrame:
    """Faits de la nouvelle société précédés de ceux de l'ancienne, publiés avant son premier dépôt."""
    if ancien is None or not len(ancien):
        return nouveau
    if nouveau is None or not len(nouveau):
        return ancien
    premier = pd.to_datetime(nouveau["depot"]).min()
    return pd.concat([ancien[pd.to_datetime(ancien["depot"]) < premier], nouveau], ignore_index=True)


def faits_societe(sec, cik: int, _profondeur: int = 0) -> pd.DataFrame:
    """Faits XBRL d'une société, avec l'historique des sociétés qu'elle prolonge s'il y a lieu."""
    faits = sec.faits(cik)
    ancien = PREDECESSEURS.get(int(cik))
    if ancien and _profondeur < 3:
        try:
            faits = fusionner_predecesseur(faits, faits_societe(sec, ancien, _profondeur + 1))
        except Exception as e:  # noqa: BLE001
            info(f"  Historique de l'ancienne société (CIK {ancien}) indisponible ({type(e).__name__}).")
    return faits


def charger_comptes(ciks: Sequence[int], sec: ClientSEC) -> Dict[int, pd.DataFrame]:
    ciks = sorted({int(c) for c in ciks if np.isfinite(c)})
    info(f"Comptes SEC : {len(ciks)} sociétés (environ {max(1, len(ciks) // 150)} min la première fois, "
         "puis cache local)…")
    tables = {}
    t0 = time.time()
    erreurs = 0
    for k, cik in enumerate(ciks):
        try:
            tables[cik] = table_fondamentale(faits_societe(sec, cik))
        except Exception as e:  # noqa: BLE001
            erreurs += 1
            info(f"  CIK {cik} ignoré ({type(e).__name__}: {str(e)[:80]})")
            tables[cik] = pd.DataFrame()
        if (k + 1) % 50 == 0:
            info(f"  comptes : {k + 1}/{len(ciks)} ({time.time() - t0:.0f} s)")
    ok = sum(1 for t in tables.values() if t is not None and not t.empty)
    info(f"Comptes SEC : {ok}/{len(ciks)} sociétés exploitables.")
    ETAT_SOURCES["sec"] = {"societes": len(ciks), "exploitables": ok, "erreurs": erreurs}
    return tables


# ===========================================================================
# 10. Screener du jour
# ===========================================================================

def debut_prix(cfg: Config) -> pd.Timestamp:
    return min(aujourdhui() - pd.Timedelta(days=430),
               pd.Timestamp(cfg.debut_backtest) - pd.Timedelta(days=430))


def seance_en_cours(index: pd.DatetimeIndex, maintenant: Optional[pd.Timestamp] = None) -> bool:
    """Vrai si la dernière date est aujourd'hui à New York et que la séance n'est pas close (16 h + marge) :
    son cours est alors provisoire."""
    if len(index) == 0:
        return False
    now = maintenant if maintenant is not None else pd.Timestamp.now(tz="America/New_York")
    return pd.Timestamp(index[-1]).date() == now.date() and (now.hour, now.minute) < (16, 30)


def _preparer_prix(prix: Dict[str, pd.DataFrame], ref: str = "SPY"):
    """Aligne tout sur les jours de cotation du S&P 500. Renvoie (rendement total, SPY,
    cours ajustés des seules divisions, facteur de divisions)."""
    adj_all = prix["adj"]
    if ref not in adj_all.columns:
        raise RuntimeError("Le S&P 500 (SPY) n'a pas pu être téléchargé : réessaie dans quelques minutes.")
    spy = adj_all[ref].dropna()
    if seance_en_cours(spy.index):
        info(f"Séance du {spy.index[-1].date()} encore ouverte à New York : son cours provisoire est ignoré.")
        spy = spy.iloc[:-1]
    garder = [c for c in adj_all.columns if c not in ("SPY", "RSP", "^IRX")]
    adj = adj_all.reindex(spy.index)[garder].ffill(limit=5)
    close_split = prix["close"].reindex(index=spy.index, columns=garder).ffill(limit=5)
    facteur = facteur_divisions(prix["splits"].reindex(columns=garder), spy.index)
    return adj, spy, close_split, facteur


def lancer_screener(cfg: Config, positions: Optional[Sequence[str]] = None) -> dict:
    """Classement du jour, signaux et portefeuille proposé."""
    t0 = time.time()
    actuel = charger_sp500_actuel(cfg)
    tickers = actuel["ticker"].tolist()
    positions = [ticker_yahoo(p) for p in (positions or []) if str(p).strip()]
    prix = telecharger_prix(set(tickers) | set(positions) | {"SPY"}, debut_prix(cfg), cfg)
    close, spy, close_split, facteur = _preparer_prix(prix)
    recent = close.iloc[-420:]
    recent_split = close_split.iloc[-420:]
    spy_r = spy.reindex(recent.index)
    mats = matrices_prix(recent, spy_r)
    i = len(recent.index) - 1
    d = recent.index[i]
    mdd = perte_max_1an(recent, [i]).iloc[0]
    infos = actuel.set_index("ticker")[["nom", "secteur", "cik"]]
    fond_d = None
    if cfg.mode == "complet":
        sec = ClientSEC(cfg)
        tables = charger_comptes(infos["cik"].dropna().astype(int).tolist(), sec)
        panneau = panneau_fondamental(tables)
        fond_d = fondamentaux_aux_dates(panneau, [d], infos["cik"].dropna().astype(int).tolist())
    revisions = lire_revisions(tickers, cfg) if cfg.revisions_analystes else None
    df = tableau_facteurs(d, i, tickers, recent, mats, mdd, infos, fond_d, revisions, cfg.mode,
                          close_split=recent_split, facteur=facteur)
    if cfg.mode == "complet":
        manque = df.index[df["capitalisation"].isna() & df["historique_ok"]].tolist()
        if manque:
            info(f"Capitalisation : {len(manque)} titres complétés via Yahoo ({', '.join(manque[:8])}"
                 f"{'…' if len(manque) > 8 else ''}).")
            metriques_marche(df, d, lire_capitalisations(manque))
    df = attribuer_signaux(calculer_scores(appliquer_filtres(df, cfg), cfg), cfg)
    df["secteur_fr"] = df["secteur"].map(SECTEURS_FR).fillna(df["secteur"])
    if "sous_secteur" in actuel:
        df["sous_secteur"] = actuel.set_index("ticker")["sous_secteur"].reindex(df.index).fillna("").astype(str)
    df = df.sort_values("score", ascending=False, na_position="last")

    gardes, achats, ventes = construire_portefeuille(df, positions, cfg,
                                                     hors_univers="Hors S&P 500 : non analysé")
    lignes = gardes + achats
    port = df.loc[lignes].copy()
    port["action"] = ["Conserver" if t in gardes else "Acheter" for t in lignes]
    port["poids"] = 1.0 / len(lignes) if lignes else np.nan
    mes = None
    if positions:
        actions_sig = {"ACHETER": "Conserver (renforcer possible)", "GARDER": "Conserver",
                       "SURVEILLER": "Conserver, sous surveillance", "VENDRE": "Vendre",
                       "DONNÉES INSUFFISANTES": "Pas d'avis (données insuffisantes)"}
        rangs = []
        for t in dict.fromkeys(positions):
            if t in df.index:
                r = df.loc[t]
                rangs.append({"ticker": t, "nom": r["nom"], "signal": r["signal"], "score": r["score"],
                              "action": actions_sig.get(r["signal"], ""), "motif": r["motif"],
                              "alertes": r["alertes"]})
            else:
                rangs.append({"ticker": t, "nom": "", "signal": "HORS UNIVERS", "score": np.nan,
                              "action": "Non analysé (hors S&P 500)", "motif": "", "alertes": ""})
        mes = pd.DataFrame(rangs)
    comptes = df["signal"].value_counts().reindex(SIGNAUX).fillna(0).astype(int)
    info(f"Classement du {d.date()} : " + ", ".join(f"{k} {v}" for k, v in comptes.items())
         + f" ({time.time() - t0:.0f} s).")
    return {"date": d, "classement": df, "portefeuille": port, "mes_positions": mes, "comptes": comptes,
            "source_univers": actuel.attrs.get("source", ""), "config": asdict(cfg),
            "cours": close_split.iloc[-506:]}   # deux ans de cours (ajustés des divisions) pour les graphiques


# ===========================================================================
# 11. Backtest
# ===========================================================================

URL_FRED = "https://fred.stlouisfed.org/graph/fredgraph.csv?id=DTB3"


def charger_taux_sans_risque(cfg: Config) -> Optional[pd.Series]:
    """Taux des bons du Trésor américain à 3 mois, en % par an (Réserve fédérale de Saint-Louis, série DTB3).
    None si la source officielle est injoignable : le backtest prend alors le taux ^IRX de Yahoo."""
    chemin = cfg.chemin("cache", "fred_dtb3.pkl")
    if _frais(chemin, 1):
        return pd.read_pickle(chemin)
    try:
        texte = http_get(URL_FRED, _ua_web(cfg), timeout=30, essais=2).text
        df = pd.read_csv(io.StringIO(texte))
        s = pd.Series(pd.to_numeric(df.iloc[:, 1], errors="coerce").values,
                      index=pd.to_datetime(df.iloc[:, 0], errors="coerce"))
        s = s[s.index.notna()].dropna().sort_index()
        if len(s) < 250:
            raise ValueError("série trop courte")
        s.to_pickle(chemin)
        return s
    except Exception as e:  # noqa: BLE001
        info(f"FRED indisponible ({type(e).__name__}) : taux sans risque de Yahoo (^IRX) à la place.")
        return pd.read_pickle(chemin) if os.path.exists(chemin) else None


def _fins_de_mois(index: pd.DatetimeIndex) -> List[pd.Timestamp]:
    s = pd.Series(index, index=index)
    return list(s.groupby([index.year, index.month]).max())


def _fin_de_mois_serie(s: pd.Series) -> pd.Series:
    s = s.dropna()
    return s.groupby([s.index.year, s.index.month]).tail(1)


def metriques(nav: pd.Series, ref: Optional[pd.Series], rf_j: Optional[pd.Series], est_ref: bool = False) -> dict:
    nav = nav.dropna()
    if len(nav) < 60:
        return {}
    r = nav.pct_change(fill_method=None).dropna()
    rf = rf_j.reindex(r.index).fillna(0.0) if rf_j is not None else pd.Series(0.0, index=r.index)
    ex = r - rf
    annees = (nav.index[-1] - nav.index[0]).days / 365.25
    cagr = (nav.iloc[-1] / nav.iloc[0]) ** (1 / annees) - 1 if annees > 0 else np.nan
    et = r.std()
    vol = et * np.sqrt(252)
    sharpe = ex.mean() / et * np.sqrt(252) if et > 0 else np.nan
    neg = np.minimum(ex, 0.0)
    dev_neg = np.sqrt((neg ** 2).mean()) * np.sqrt(252)
    sortino = ex.mean() * 252 / dev_neg if dev_neg > 0 else np.nan
    mdd = (nav / nav.cummax() - 1).min()
    mens = _fin_de_mois_serie(nav)
    r12 = (mens / mens.shift(12) - 1).dropna()
    res = {"Rendement annualisé": cagr, "Volatilité annuelle": vol, "Ratio de Sharpe": sharpe,
           "Ratio de Sortino": sortino, "Perte maximale": mdd,
           "Ratio de Calmar": cagr / abs(mdd) if mdd < 0 else np.nan,
           "Pire période de 12 mois": r12.min() if len(r12) else np.nan,
           "Périodes de 12 mois en perte": (r12 < 0).mean() if len(r12) else np.nan}
    if ref is not None and not est_ref:
        refm = _fin_de_mois_serie(ref.reindex(nav.index).ffill())
        r12r = (refm / refm.shift(12) - 1).reindex(r12.index)
        res["Périodes de 12 mois battant le S&P 500"] = (r12 > r12r).mean() if len(r12) else np.nan
        rr = ref.reindex(nav.index).pct_change(fill_method=None).reindex(r.index)
        ok = rr.notna()
        if ok.sum() > 60:
            res["Bêta vs S&P 500"] = float(np.cov(r[ok], rr[ok])[0, 1] / rr[ok].var())
    else:
        res["Périodes de 12 mois battant le S&P 500"] = np.nan
        res["Bêta vs S&P 500"] = 1.0 if est_ref else np.nan
    return res


def rendements_annuels(navs: pd.DataFrame) -> pd.DataFrame:
    out = {}
    for c in navs.columns:
        s = navs[c].dropna()
        if s.empty:
            continue
        fins = s.groupby(s.index.year).last()
        debuts = fins.shift(1)
        debuts.iloc[0] = s.iloc[0]
        out[c] = fins / debuts - 1
    return pd.DataFrame(out)


def _spearman(a: pd.Series, b: pd.Series) -> float:
    return float(a.rank().corr(b.rank()))


def _tstat(s: pd.Series) -> float:
    s = s.dropna()
    if len(s) < 3 or s.std(ddof=1) == 0:
        return np.nan
    return float(s.mean() / s.std(ddof=1) * np.sqrt(len(s)))


FACTEURS_ANALYSE = [("Score global", "score"), ("Potentiel", "potentiel_pct"), ("Robustesse", "robustesse_pct"),
                    ("Momentum 12-1", "p_momentum"), ("Résultats (surprises)", "p_resultats"),
                    ("Valorisation", "p_valeur"), ("Qualité", "p_qualite"), ("Faible risque", "p_faible_risque")]


def analyse_facteurs(h: pd.DataFrame) -> Dict[str, pd.DataFrame]:
    lignes = []
    groupes = list(h.groupby("date"))
    for nom, col in FACTEURS_ANALYSE:
        if col not in h or h[col].notna().sum() == 0:
            continue
        ic1, ic12 = {}, {}
        for d, g in groupes:
            for horizon, cible in (("r1m", ic1), ("r12m", ic12)):
                x = g[[col, horizon]].dropna()
                if len(x) >= 30:
                    cible[d] = _spearman(x[col], x[horizon])
        s1, s12 = pd.Series(ic1, dtype=float), pd.Series(ic12, dtype=float)
        lignes.append({"Facteur": nom, "IC moyen à 1 mois": s1.mean(), "t-stat (1 mois)": _tstat(s1),
                       "Mois avec IC > 0": (s1 > 0).mean() if len(s1) else np.nan,
                       "IC moyen à 12 mois": s12.mean(),
                       "t-stat (12 mois, sans chevauchement)": _tstat(s12.iloc[::12])})
    ic = pd.DataFrame(lignes)
    q = h.dropna(subset=["score", "r12m"]).copy()
    if q.empty:
        return {"ic": ic, "quintiles": pd.DataFrame(), "signaux": pd.DataFrame()}
    q["quintile"] = q.groupby("date")["score"].transform(lambda s: np.ceil(s.rank(pct=True) * 5).clip(1, 5))
    agg = q.groupby(["date", "quintile"]).agg(r=("r12m", "mean"), neg=("r12m", lambda x: (x < 0).mean()),
                                              vol=("vol", "mean")).reset_index()
    quint = agg.groupby("quintile").agg(r=("r", "mean"), neg=("neg", "mean"), vol=("vol", "mean"))
    quint.index = [f"Q{int(i)}" + (" (scores faibles)" if i == 1 else " (scores élevés)" if i == 5 else "")
                   for i in quint.index]
    quint.columns = ["Rendement moyen à 12 mois", "Part des titres en baisse à 12 mois", "Volatilité moyenne"]
    sg = q.groupby(["date", "signal"]).agg(r=("r12m", "mean"), neg=("r12m", lambda x: (x < 0).mean()),
                                           n=("r12m", "size")).reset_index()
    sig = sg.groupby("signal").agg(r=("r", "mean"), neg=("neg", "mean"), n=("n", "mean"))
    sig = sig.reindex([s for s in SIGNAUX if s in sig.index])
    sig.columns = ["Rendement moyen à 12 mois", "Part des titres en baisse à 12 mois", "Titres par mois"]
    return {"ic": ic, "quintiles": quint, "signaux": sig}


def lancer_backtest(cfg: Config) -> dict:
    """Rejoue la stratégie mois par mois sur la composition historique de l'indice."""
    t0 = time.time()
    actuel = charger_sp500_actuel(cfg)
    hist = charger_historique(cfg)
    compo = Composition(hist, actuel)
    debut = pd.Timestamp(cfg.debut_backtest)
    fin = pd.Timestamp(cfg.fin_backtest) if cfg.fin_backtest else aujourdhui()
    univers = compo.tous_depuis(debut, fin)
    info(f"Backtest {debut.date()} → {fin.date()} : {len(univers)} titres ont fait partie de l'indice.")
    prix = telecharger_prix(univers | {"SPY", "RSP", "^IRX"}, debut - pd.Timedelta(days=430), cfg)
    irx = prix["adj"]["^IRX"].dropna() if "^IRX" in prix["adj"] else None
    rsp = prix["adj"]["RSP"].dropna() if "RSP" in prix["adj"] else None
    close, spy, close_split, facteur = _preparer_prix(prix)
    close = close.loc[close.index <= fin]
    close_split = close_split.loc[close_split.index <= fin]
    spy = spy.loc[spy.index <= fin]
    taux = charger_taux_sans_risque(cfg)
    ETAT_SOURCES["taux"] = {"source": "FRED (DTB3)" if taux is not None else ("Yahoo (^IRX)" if irx is not None
                                                                             else "aucune (taux nul)")}
    if taux is None:
        taux = irx
    rf_j = (taux.reindex(close.index.union(taux.index)).ffill().reindex(close.index) / 100 / 252
            ) if taux is not None else None

    sec = ClientSEC(cfg) if cfg.mode == "complet" or cfg.email_valide() else None
    infos = infos_tickers(list(close.columns), actuel, sec)
    dates = [d for d in _fins_de_mois(close.index) if d >= debut and d < close.index[-1]]
    if len(dates) < 13:
        raise RuntimeError("Période de backtest trop courte (moins d'un an).")
    pos = {d: close.index.get_loc(d) for d in dates}
    fond = {}
    if cfg.mode == "complet":
        ciks = infos["cik"].dropna().astype(int).tolist()
        tables = charger_comptes(ciks, sec)
        panneau = panneau_fondamental(tables)
        tout = fondamentaux_aux_dates(panneau, dates, ciks)
        if not tout.empty:
            fond = {pd.Timestamp(d): g for d, g in tout.groupby("date")}
    info("Calcul des facteurs de prix…")
    mats = matrices_prix(close, spy)
    mdd = perte_max_1an(close, [pos[d] for d in dates])
    R = mats["ret"].fillna(0.0)
    n_jours = len(close.index)

    nav = pd.Series(np.nan, index=close.index)
    ew = pd.Series(np.nan, index=close.index)
    j0 = pos[dates[0]]
    nav.iloc[j0], ew.iloc[j0] = 1.0, 1.0
    positions: Dict[str, float] = {}
    cash, v_ew, j_dernier = 1.0, 1.0, j0
    journal, histo, couverture = [], [], {}

    def avancer(j_a: int, j_b: int):
        nonlocal positions
        if j_b <= j_a:
            return
        if positions:
            tick = list(positions)
            g = (1.0 + R.iloc[j_a + 1: j_b + 1][tick]).cumprod().values
            vals = g * np.array([positions[t] for t in tick])[None, :]
            nav.iloc[j_a + 1: j_b + 1] = vals.sum(axis=1) + cash
            positions = {t: float(vals[-1, c]) for c, t in enumerate(tick)}
        else:
            nav.iloc[j_a + 1: j_b + 1] = cash

    info(f"Simulation de {len(dates)} mois…")
    for k, d in enumerate(dates):
        i = pos[d]
        brut = compo.brut(d)
        membres = compo.membres(d)
        dispo = sorted(m for m in membres if m in close.columns)
        df = tableau_facteurs(d, i, dispo, close, mats, mdd.loc[d], infos, fond.get(pd.Timestamp(d)),
                              None, cfg.mode, close_split=close_split, facteur=facteur)
        df = attribuer_signaux(calculer_scores(appliquer_filtres(df, cfg), cfg), cfg)
        valides = df.index[df["historique_ok"]].tolist()
        couverture[d] = len(valides) / max(1, len(brut))

        ex = min(i + 1, n_jours - 1)                       # exécution à la clôture suivante
        j_suiv = min(pos[dates[k + 1]] + 1, n_jours - 1) if k + 1 < len(dates) else n_jours - 1
        avancer(j_dernier, ex)
        j_dernier = ex
        valeur = sum(positions.values()) + cash
        rebal = (cfg.rebalancement == "mensuel") or (d.month in (3, 6, 9, 12)) or k == 0
        rotation, cout, achats_k, ventes_k = 0.0, 0.0, [], []
        if rebal:
            gardes, achats_k, ventes = construire_portefeuille(df, list(positions), cfg)
            ventes_k = [t for t, _ in ventes]
            cible = gardes + achats_k
            w = 1.0 / len(cible) if cible else 0.0
            union = set(positions) | set(cible)
            rotation = sum(abs((w * valeur if t in cible else 0.0) - positions.get(t, 0.0)) for t in union)
            cout = cfg.frais * rotation
            nette = valeur - cout
            positions = {t: w * nette for t in cible}
            cash = 0.0 if cible else nette
        elif cfg.ventes_urgence and positions:
            for t in list(positions):
                if t not in df.index or bool(df.at[t, "echec_filtre"]):
                    v = positions.pop(t)
                    rotation += v
                    cout += v * cfg.frais
                    cash += v * (1 - cfg.frais)
                    ventes_k.append(t)
        journal.append({"date": d, "rebalancement": rebal, "positions": len(positions),
                        "titres": len(df), "acheter": int((df["signal"] == "ACHETER").sum()),
                        "insuffisantes": int((df["signal"] == "DONNÉES INSUFFISANTES").sum()),
                        "achats": ", ".join(achats_k), "ventes": ", ".join(ventes_k),
                        "rotation": rotation / valeur if valeur > 0 else 0.0,
                        "frais": cout / valeur if valeur > 0 else 0.0,
                        "liquidites": cash / max(valeur - cout, 1e-12),
                        "portefeuille": ", ".join(sorted(positions))})
        # Univers équipondéré (mêmes données, rebalancé chaque mois, sans frais)
        if valides and j_suiv > ex:
            g = (1.0 + R.iloc[ex + 1: j_suiv + 1][valides]).cumprod().mean(axis=1).values * v_ew
            ew.iloc[ex + 1: j_suiv + 1] = g
            v_ew = float(g[-1])
        cols = ["score", "potentiel_pct", "robustesse_pct", "p_momentum", "p_resultats", "p_valeur",
                "p_qualite", "p_faible_risque", "signal", "vol"]
        h = df[[c for c in cols if c in df]].copy()
        h["date"] = d
        histo.append(h)
        if (k + 1) % 24 == 0:
            info(f"  {d.date()} : {k + 1}/{len(dates)} mois simulés")
    avancer(j_dernier, n_jours - 1)
    ew.iloc[j0: pos[dates[0]] + 2] = ew.iloc[j0: pos[dates[0]] + 2].fillna(1.0)
    ew = ew.ffill()

    # Rendements futurs pour l'analyse des facteurs
    prix_mois = close.loc[dates]
    r1 = prix_mois.shift(-1) / prix_mois - 1
    r12 = prix_mois.shift(-12) / prix_mois - 1
    hh = pd.concat(histo).reset_index().rename(columns={"index": "ticker"})
    hh["r1m"] = [r1.at[d, t] if t in r1.columns else np.nan for d, t in zip(hh["date"], hh["ticker"])]
    hh["r12m"] = [r12.at[d, t] if t in r12.columns else np.nan for d, t in zip(hh["date"], hh["ticker"])]
    analyse = analyse_facteurs(hh)

    debut_nav = close.index[j0]
    navs = pd.DataFrame({"Stratégie": nav, "S&P 500 (SPY)": spy / spy.loc[debut_nav],
                         "Univers équipondéré": ew}).loc[debut_nav:]
    if rsp is not None and rsp.index[0] <= debut_nav:
        rr = rsp.reindex(close.index).ffill()
        navs["S&P 500 équipondéré (RSP)"] = (rr / rr.loc[debut_nav]).loc[debut_nav:]
    navs = navs.ffill()
    periodes = [("Période complète", navs.index[0], navs.index[-1]),
                (f"Construction ({navs.index[0].year}-{pd.Timestamp(cfg.debut_validation).year - 1})",
                 navs.index[0], pd.Timestamp(cfg.debut_validation) - pd.Timedelta(days=1)),
                (f"Validation ({pd.Timestamp(cfg.debut_validation).year}-{navs.index[-1].year})",
                 pd.Timestamp(cfg.debut_validation), navs.index[-1])]
    tables_m = {}
    for nom, a, b in periodes:
        sous = navs.loc[a:b]
        if len(sous) < 120:
            continue
        tables_m[nom] = pd.DataFrame({c: metriques(sous[c], sous["S&P 500 (SPY)"], rf_j,
                                                   est_ref=(c == "S&P 500 (SPY)")) for c in sous.columns})
    jr = pd.DataFrame(journal)
    reb = jr[jr["rebalancement"]]
    annees = (navs.index[-1] - navs.index[0]).days / 365.25
    synthese = {
        "Rotation annuelle moyenne": jr["rotation"].sum() / annees if annees > 0 else np.nan,
        "Frais cumulés (% du capital)": jr["frais"].sum(),
        "Lignes en moyenne": reb["positions"].mean() if len(reb) else np.nan,
        "Liquidités moyennes": jr["liquidites"].mean(),
        "Couverture des données (moyenne)": float(np.mean(list(couverture.values()))),
    }
    info(f"Backtest terminé en {(time.time() - t0) / 60:.1f} min.")
    for nom, tab in tables_m.items():
        def valeurs(ind: str) -> str:
            return ", ".join(f"{c} {_fr(tab.at[ind, c] * 100)} %" for c in tab.columns
                             if ind in tab.index and np.isfinite(tab.at[ind, c]))
        info(f"  {nom} : rendement annualisé {valeurs('Rendement annualisé')} ; "
             f"perte maximale {valeurs('Perte maximale')}.")
    return {"navs": navs, "metriques": tables_m, "annuel": rendements_annuels(navs), "journal": jr,
            "analyse": analyse, "couverture": pd.Series(couverture), "synthese": synthese,
            "config": asdict(cfg), "debut": navs.index[0], "fin": navs.index[-1],
            "historique_scores": hh[["date", "ticker", "score", "signal"]].copy()}


# ===========================================================================
# 12. Rapports : Excel et page HTML
# ===========================================================================

SURFACE, ENCRE, ENCRE2, MUET, GRILLE, AXE = "#fcfcfb", "#0b0b0b", "#52514e", "#898781", "#e1e0d9", "#c3c2b7"
COULEURS = {"Stratégie": "#2a78d6", "S&P 500 (SPY)": "#eb6834", "Univers équipondéré": "#1baf7a",
            "S&P 500 équipondéré (RSP)": "#6250d6"}

COLONNES_AFFICHAGE = [
    ("nom", "Société", "txt"), ("secteur_fr", "Secteur", "txt"), ("signal", "Signal", "sig"),
    ("score", "Score", "n0"), ("potentiel_pct", "Potentiel", "n0"), ("robustesse_pct", "Robustesse", "n0"),
    ("mom", "Momentum 12-1", "pct"), ("revision", "Révision BPA 90 j", "pct"),
    ("surprise", "Surprise résultats (% actif)", "pct2"), ("rend_fcf", "Rendement FCF", "pct"),
    ("ebit_ve", "EBIT / VE", "pct"), ("rend_benef", "Rendement bénéfices", "pct"),
    ("rentab_brute", "Rentabilité brute", "pct"), ("roa_moy", "ROA", "pct"), ("fscore_9", "F-score (/9)", "n1"),
    ("couverture", "Couverture intérêts", "n1"), ("dette_ebitda", "Dette nette / EBITDA", "n1"),
    ("accruals", "Accruals", "pct"), ("beneish", "M-score Beneish", "n2"), ("vol", "Volatilité 1 an", "pct"),
    ("beta", "Bêta 1 an", "n2"), ("mdd", "Perte max 1 an", "pct"), ("capi_md", "Capitalisation (Md$)", "n1"),
    ("motif", "Motif", "txt"), ("alertes", "Alertes", "txt"),
]
FORMATS_XL = {"n0": "0", "n1": "0.0", "n2": "0.00", "pct": "0.0%", "pct2": "0.00%"}


def tableau_affichage(df: pd.DataFrame, colonnes=None) -> pd.DataFrame:
    d = df.copy()
    if "capitalisation" in d:
        d["capi_md"] = d["capitalisation"] / 1e9
    cols = colonnes or COLONNES_AFFICHAGE
    out = pd.DataFrame(index=d.index)
    out.index.name = "Ticker"
    for c, titre, _ in cols:
        if c in d:
            v = d[c]
            if pd.api.types.is_numeric_dtype(v):
                v = v.astype(float).where(np.isfinite(v.astype(float)))
            out[titre] = v
    return out


def _figure(largeur: float, hauteur: float, ncols: int = 1):
    """Figure matplotlib autonome (sans pyplot) : aucun effet sur l'affichage du notebook."""
    from matplotlib.figure import Figure

    fig = Figure(figsize=(largeur, hauteur))
    return fig, fig.subplots(1, ncols)


def _fig_b64(fig) -> str:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight", facecolor=SURFACE)
    return base64.b64encode(buf.getvalue()).decode("ascii")


def _axes(ax, titre_y: str = ""):
    ax.set_facecolor(SURFACE)
    for c in ("top", "right"):
        ax.spines[c].set_visible(False)
    for c in ("left", "bottom"):
        ax.spines[c].set_color(AXE)
    ax.tick_params(colors=MUET, labelsize=9)
    ax.grid(axis="y", color=GRILLE, linewidth=0.8)
    ax.set_axisbelow(True)
    if titre_y:
        ax.set_ylabel(titre_y, color=ENCRE2, fontsize=9)


def _pct_fmt(ax, axe="y"):
    from matplotlib.ticker import FuncFormatter

    f = FuncFormatter(lambda v, _: f"{round(v * 100, 2):g} %".replace(".", ","))
    (ax.yaxis if axe == "y" else ax.xaxis).set_major_formatter(f)


def _etiquettes_annees(index, debut, fin) -> List[str]:
    out = []
    for a in index:
        partiel = (int(a) == pd.Timestamp(fin).year and pd.Timestamp(fin).month < 12) or \
                  (int(a) == pd.Timestamp(debut).year and pd.Timestamp(debut).month > 1)
        out.append(f"{a}*" if partiel else str(a))
    return out


def graphiques_backtest(bt: dict) -> Dict[str, str]:
    from matplotlib.ticker import FuncFormatter

    navs = bt["navs"]
    out = {}
    series = [c for c in ["Stratégie", "S&P 500 (SPY)", "Univers équipondéré"] if c in navs]

    fig, ax = _figure(9, 4.2)
    _axes(ax, "Valeur (base 1, échelle log)")
    for c in series:
        ax.plot(navs.index, navs[c], color=COULEURS[c], lw=1.8, label=c, zorder=3 if c == "Stratégie" else 2)
        ax.annotate(f"{navs[c].iloc[-1]:.2f}".replace(".", ","), (navs.index[-1], navs[c].iloc[-1]),
                    xytext=(6, 0), textcoords="offset points", va="center", fontsize=9, color=ENCRE2)
    ax.set_yscale("log")
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:g}".replace(".", ",")))
    ax.legend(frameon=False, fontsize=9, loc="upper left", labelcolor=ENCRE)
    out["valeur"] = _fig_b64(fig)

    fig, ax = _figure(9, 3.4)
    _axes(ax, "Baisse depuis le plus haut")
    dd_s = navs["Stratégie"] / navs["Stratégie"].cummax() - 1
    dd_m = navs["S&P 500 (SPY)"] / navs["S&P 500 (SPY)"].cummax() - 1
    pire = lambda s: f"{s.min() * 100:.1f} %".replace(".", ",")   # noqa: E731
    ax.fill_between(navs.index, dd_s.values, 0, color=COULEURS["Stratégie"], alpha=0.18, linewidth=0)
    ax.plot(navs.index, dd_s, color=COULEURS["Stratégie"], lw=1.4, zorder=3,
            label=f"Stratégie (pire baisse {pire(dd_s)})")
    ax.plot(navs.index, dd_m, color=COULEURS["S&P 500 (SPY)"], lw=1.1, zorder=2,
            label=f"S&P 500 (pire baisse {pire(dd_m)})")
    _pct_fmt(ax)
    ax.legend(frameon=False, fontsize=9, loc="lower left", bbox_to_anchor=(0, 1.0), ncol=2, labelcolor=ENCRE)
    out["baisse"] = _fig_b64(fig)

    an = bt["annuel"][series]
    fig, ax = _figure(9, 3.8)
    _axes(ax, "Rendement de l'année")
    x = np.arange(len(an.index))
    larg = 0.8 / len(series)
    for k, c in enumerate(series):
        ax.bar(x + (k - (len(series) - 1) / 2) * larg, an[c].values, width=larg, color=COULEURS[c],
               label=c, edgecolor=SURFACE, linewidth=1.5)
    ax.axhline(0, color=AXE, lw=1)
    ax.set_xticks(x)
    ax.set_xticklabels(_etiquettes_annees(an.index, bt["debut"], bt["fin"]), fontsize=8)
    _pct_fmt(ax)
    ax.legend(frameon=False, fontsize=9, loc="upper left", ncol=3, labelcolor=ENCRE)
    out["annuel"] = _fig_b64(fig)

    q = bt["analyse"].get("quintiles")
    if q is not None and not q.empty:
        fig, axs = _figure(9, 3.4, ncols=2)
        for ax, col, titre in zip(axs, ["Rendement moyen à 12 mois", "Part des titres en baisse à 12 mois"],
                                  ["Rendement moyen à 12 mois", "Part des titres en baisse à 12 mois"]):
            _axes(ax)
            etiquettes = [i.split(" ")[0] for i in q.index]
            barres = ax.bar(etiquettes, q[col].values, color="#2a78d6", edgecolor=SURFACE, linewidth=1.5)
            for b, v in zip(barres, q[col].values):
                ax.annotate(f"{v * 100:.1f} %".replace(".", ","), (b.get_x() + b.get_width() / 2, v),
                            xytext=(0, 3), textcoords="offset points", ha="center", fontsize=8, color=ENCRE2)
            ax.set_title(titre, fontsize=10, color=ENCRE, loc="left")
            _pct_fmt(ax)
            ax.set_xlabel("Quintile du score (Q1 = faible, Q5 = élevé)", fontsize=8, color=MUET)
        fig.tight_layout()
        out["quintiles"] = _fig_b64(fig)

    cv = bt["couverture"]
    fig, ax = _figure(9, 2.6)
    _axes(ax, "Membres de l'indice avec données")
    ax.plot(cv.index, cv.values, color="#2a78d6", lw=1.6)
    ax.set_ylim(0, 1.02)
    _pct_fmt(ax)
    out["couverture"] = _fig_b64(fig)
    return out


def graphique_secteurs(port: pd.DataFrame) -> Optional[str]:
    if port is None or port.empty:
        return None
    c = port["secteur_fr"].value_counts().sort_values()
    fig, ax = _figure(7, 0.35 * len(c) + 0.8)
    _axes(ax)
    ax.grid(axis="y", visible=False)
    ax.grid(axis="x", color=GRILLE, linewidth=0.8)
    barres = ax.barh(c.index, c.values, color="#2a78d6", edgecolor=SURFACE, linewidth=1.5)
    for b, v in zip(barres, c.values):
        ax.annotate(str(v), (v, b.get_y() + b.get_height() / 2), xytext=(4, 0), textcoords="offset points",
                    va="center", fontsize=9, color=ENCRE2)
    from matplotlib.ticker import MaxNLocator

    ax.xaxis.set_major_locator(MaxNLocator(integer=True))
    ax.tick_params(axis="y", colors=ENCRE, labelsize=9)
    ax.set_xlabel("Nombre de lignes", fontsize=9, color=ENCRE2)
    return _fig_b64(fig)


def _fmt(v, kind: str) -> str:
    if v is None or (isinstance(v, float) and not np.isfinite(v)) or (not isinstance(v, str) and pd.isna(v)):
        return "—"
    try:
        if kind == "pct":
            return f"{float(v) * 100:.1f} %".replace(".", ",")
        if kind == "pct2":
            return f"{float(v) * 100:.2f} %".replace(".", ",")
        if kind in ("n0", "n1", "n2"):
            return f"{float(v):.{int(kind[1])}f}".replace(".", ",")
    except (TypeError, ValueError):
        pass
    return html.escape(str(v))


PASTILLES = {"ACHETER": ("▲", "achat"), "GARDER": ("●", "garder"), "SURVEILLER": ("◆", "surv"),
             "VENDRE": ("▼", "vente"), "DONNÉES INSUFFISANTES": ("?", "nd"), "HORS UNIVERS": ("–", "nd")}


def _pastille(sig: str) -> str:
    sym, cls = PASTILLES.get(str(sig), ("", "nd"))
    return f'<span class="sig {cls}">{sym} {html.escape(str(sig))}</span>'


def _table_html(df: pd.DataFrame, colonnes: List[Tuple[str, str, str]], index_titre: Optional[str] = None) -> str:
    tete = ([f'<th class="g">{html.escape(index_titre)}</th>'] if index_titre else []) + [
        f'<th class="{"g" if k in ("txt", "sig") else ""}">{html.escape(t)}</th>' for _, t, k in colonnes]
    corps = []
    for idx, ligne in df.iterrows():
        cellules = [f'<td class="g"><b>{html.escape(str(idx))}</b></td>'] if index_titre else []
        for c, _, k in colonnes:
            v = ligne.get(c, None)
            if k == "sig":
                cellules.append(f'<td class="g">{_pastille(v)}</td>')
            elif k == "txt":
                cellules.append(f'<td class="g texte">{_fmt(v, k)}</td>')
            else:
                cellules.append(f"<td>{_fmt(v, k)}</td>")
        corps.append("<tr>" + "".join(cellules) + "</tr>")
    return ('<div class="tablewrap"><table><thead><tr>' + "".join(tete) + "</tr></thead><tbody>"
            + "".join(corps) + "</tbody></table></div>")


CSS = """
:root{--fond:#f9f9f7;--carte:#fcfcfb;--encre:#0b0b0b;--encre2:#52514e;--muet:#898781;--trait:#e1e0d9;
--bord:rgba(11,11,11,.10);--accent:#2a78d6}
@media (prefers-color-scheme: dark){:root{--fond:#0d0d0d;--carte:#1a1a19;--encre:#f0efec;--encre2:#c3c2b7;
--muet:#898781;--trait:#2c2c2a;--bord:rgba(255,255,255,.10);--accent:#3987e5}
figure{background:#fcfcfb!important}figcaption{color:#52514e!important}}
*{box-sizing:border-box}body{margin:0;background:var(--fond);color:var(--encre);
font:15px/1.55 system-ui,-apple-system,"Segoe UI",sans-serif}
main{max-width:1180px;margin:0 auto;padding:28px 16px 72px}
h1{font-size:27px;margin:0 0 4px;letter-spacing:-.01em}h2{font-size:19px;margin:38px 0 10px}
h3{font-size:15px;margin:22px 0 8px;color:var(--encre2)}
.sous{color:var(--encre2);margin:0 0 20px}.note{color:var(--encre2);font-size:13.5px;max-width:860px}
.cartes{display:flex;flex-wrap:wrap;gap:10px;margin:14px 0}
.carte{flex:1 1 150px;background:var(--carte);border:1px solid var(--bord);border-radius:10px;padding:12px 14px}
.carte .val{font-size:24px;font-weight:650}.carte .lib{color:var(--encre2);font-size:13px}
.tablewrap{overflow-x:auto;border:1px solid var(--bord);border-radius:10px;background:var(--carte)}
table{border-collapse:collapse;width:100%;font-variant-numeric:tabular-nums;font-size:13px}
th,td{padding:6px 9px;border-bottom:1px solid var(--trait);text-align:right;white-space:nowrap}
th.g,td.g{text-align:left}td.texte{white-space:normal;min-width:260px;max-width:460px;color:var(--encre2);
font-size:12.5px}figure.petit{max-width:720px}
th{color:var(--encre2);font-weight:600;background:var(--carte);position:sticky;top:0}
tbody tr:last-child td{border-bottom:0}
.sig{font-weight:650;padding:2px 9px;border-radius:999px;font-size:12px;display:inline-block;white-space:nowrap}
.sig.achat{background:#e3f4e3;color:#006300}.sig.garder{background:#e3eefb;color:#1c5cab}
.sig.surv{background:#fdf1d8;color:#7a5200}.sig.vente{background:#fbe4e4;color:#a12b2b}
.sig.nd{background:#eeede9;color:#52514e}
figure{margin:12px 0 0;background:var(--carte);border:1px solid var(--bord);border-radius:10px;padding:10px}
figure img{width:100%;height:auto;display:block}figcaption{color:var(--encre2);font-size:13px;margin-top:6px}
.avert{border-left:3px solid #fab219;padding:10px 14px;background:var(--carte);border-radius:0 10px 10px 0;
color:var(--encre2);font-size:13.5px;margin:16px 0}
.grille2{display:grid;grid-template-columns:repeat(auto-fit,minmax(340px,1fr));gap:14px}
ul{padding-left:20px}li{margin:3px 0}
"""


def _cartes(items: List[Tuple[str, str]]) -> str:
    return '<div class="cartes">' + "".join(
        f'<div class="carte"><div class="val">{html.escape(v)}</div><div class="lib">{html.escape(l)}</div></div>'
        for l, v in items) + "</div>"


def _metriques_html(tab: pd.DataFrame) -> str:
    formats = {"Rendement annualisé": "pct", "Volatilité annuelle": "pct", "Ratio de Sharpe": "n2",
               "Ratio de Sortino": "n2", "Perte maximale": "pct", "Ratio de Calmar": "n2",
               "Pire période de 12 mois": "pct", "Périodes de 12 mois en perte": "pct",
               "Périodes de 12 mois battant le S&P 500": "pct", "Bêta vs S&P 500": "n2"}
    lignes = []
    for m in tab.index:
        k = formats.get(m, "n2")
        lignes.append("<tr><td class='g'>" + html.escape(m) + "</td>"
                      + "".join(f"<td>{_fmt(tab.at[m, c], k)}</td>" for c in tab.columns) + "</tr>")
    tete = "<th class='g'>Indicateur</th>" + "".join(f"<th>{html.escape(c)}</th>" for c in tab.columns)
    return f'<div class="tablewrap"><table><thead><tr>{tete}</tr></thead><tbody>{"".join(lignes)}</tbody></table></div>'


def section_backtest_html(bt: dict) -> str:
    g = graphiques_backtest(bt)
    s = bt["synthese"]
    cfg = bt["config"]
    p = []
    p.append("<h2>Backtest : la stratégie rejouée mois par mois</h2>")
    p.append(f"<p class='note'>Du {bt['debut']:%d/%m/%Y} au {bt['fin']:%d/%m/%Y}, sur la composition historique "
             f"de l'indice, {cfg['n_positions']} lignes à poids égal, rebalancement {cfg['rebalancement']}, "
             f"frais de {_fr(cfg['frais'] * 100, 2)} % par transaction, exécution à la clôture du lendemain "
             "du signal. Comparaison clé : <b>l'univers équipondéré</b>, construit avec exactement les mêmes "
             "données (et donc les mêmes trous) que la stratégie."
             + (" <b>Mode « prix seulement »</b> : momentum et faible risque, sans les comptes SEC."
                if cfg.get("mode") == "prix" else "") + "</p>")
    p.append(_cartes([("Rotation annuelle (achats + ventes)", _fmt(s["Rotation annuelle moyenne"], "pct")),
                      ("Frais cumulés", _fmt(s["Frais cumulés (% du capital)"], "pct")),
                      ("Lignes en moyenne", _fmt(s["Lignes en moyenne"], "n1")),
                      ("Données disponibles", _fmt(s["Couverture des données (moyenne)"], "pct"))]))
    p.append(f'<figure><img alt="Valeur du portefeuille" src="data:image/png;base64,{g["valeur"]}">'
             "<figcaption>Valeur d'un portefeuille de base 1 au départ, dividendes réinvestis, "
             "échelle logarithmique.</figcaption></figure>")
    p.append(f'<figure><img alt="Baisses" src="data:image/png;base64,{g["baisse"]}">'
             "<figcaption>Baisse depuis le dernier plus haut : c'est la mesure du risque de perte.</figcaption></figure>")
    for nom, tab in bt["metriques"].items():
        p.append(f"<h3>{html.escape(nom)}</h3>" + _metriques_html(tab))
    p.append(f'<figure><img alt="Rendements annuels" src="data:image/png;base64,{g["annuel"]}"></figure>')
    an = bt["annuel"].copy()
    an.index = _etiquettes_annees(an.index, bt["debut"], bt["fin"])
    p.append("<h3>Rendements par année</h3>" + _table_html(an, [(c, c, "pct") for c in an.columns], "Année")
             + "<p class='note'>* année incomplète.</p>")
    a = bt["analyse"]
    if "quintiles" in g:
        p.append("<h2>Le score prédit-il quelque chose ?</h2>")
        p.append("<p class='note'>Chaque mois, les titres sont rangés en 5 groupes selon leur score. On regarde "
                 "ensuite ce qu'ils ont fait dans les 12 mois suivants. Un score utile donne des rendements qui "
                 "montent de Q1 à Q5 et une part de titres en baisse qui diminue.</p>")
        p.append(f'<figure><img alt="Quintiles" src="data:image/png;base64,{g["quintiles"]}"></figure>')
        p.append(_table_html(a["quintiles"], [(c, c, "pct") for c in a["quintiles"].columns], "Quintile"))
    if a.get("signaux") is not None and not a["signaux"].empty:
        p.append("<h3>Ce que sont devenus les titres selon leur signal (12 mois plus tard)</h3>")
        sg = a["signaux"].copy()
        p.append(_table_html(sg, [("Rendement moyen à 12 mois", "Rendement moyen à 12 mois", "pct"),
                                  ("Part des titres en baisse à 12 mois", "Part en baisse à 12 mois", "pct"),
                                  ("Titres par mois", "Titres par mois", "n0")], "Signal"))
    if a.get("ic") is not None and not a["ic"].empty:
        p.append("<h3>Pouvoir prédictif de chaque famille de facteurs</h3>")
        p.append("<p class='note'>IC = corrélation de rang entre le facteur et le rendement futur, mois par mois. "
                 "En actions, un IC moyen de 0,02 à 0,05 est déjà intéressant ; une t-stat au-dessus de 2 "
                 "(idéalement 3) indique un effet probablement réel.</p>")
        ic = a["ic"].set_index("Facteur")
        p.append(_table_html(ic, [("IC moyen à 1 mois", "IC moyen 1 mois", "n2"), ("t-stat (1 mois)", "t-stat", "n1"),
                                  ("Mois avec IC > 0", "Mois avec IC > 0", "pct"),
                                  ("IC moyen à 12 mois", "IC moyen 12 mois", "n2"),
                                  ("t-stat (12 mois, sans chevauchement)", "t-stat 12 mois", "n1")], "Facteur"))
    p.append(f'<figure><img alt="Couverture" src="data:image/png;base64,{g["couverture"]}">'
             "<figcaption>Part des membres de l'indice pour lesquels Yahoo fournit des cours. Le reste "
             "(sociétés rachetées ou disparues) manque au backtest.</figcaption></figure>")
    p.append("<div class='avert'><b>À lire avant d'y croire.</b> Un backtest décrit le passé, pas l'avenir. "
             "Les sociétés disparues manquent (biais favorable), les impôts ne sont pas déduits, et chaque "
             "réglage modifié après avoir vu ces résultats rend le prochain backtest plus optimiste qu'il ne "
             "devrait. Seule la période de validation, jamais retouchée, donne une idée honnête.</div>")
    return "\n".join(p)


def rapport_html(res: Optional[dict], bt: Optional[dict]) -> str:
    p = []
    date = res["date"] if res else (bt["fin"] if bt else aujourdhui())
    p.append(f"<h1>Screener S&amp;P 500</h1><p class='sous'>Données au {pd.Timestamp(date):%d/%m/%Y} · "
             f"version {VERSION} · outil d'aide à la décision, pas un conseil en investissement.</p>")
    if res:
        c = res["comptes"]
        cfg = res["config"]
        p.append(_cartes([("Titres analysés", str(int(c.sum()))), ("ACHETER", str(c.get("ACHETER", 0))),
                          ("GARDER", str(c.get("GARDER", 0))), ("SURVEILLER", str(c.get("SURVEILLER", 0))),
                          ("VENDRE", str(c.get("VENDRE", 0))),
                          ("Données insuffisantes", str(c.get("DONNÉES INSUFFISANTES", 0)))]))
        mes = res.get("mes_positions")
        if mes is not None and not mes.empty:
            p.append("<h2>Tes positions</h2>")
            p.append(_table_html(mes.set_index("ticker"),
                                 [("nom", "Société", "txt"), ("signal", "Signal", "sig"), ("score", "Score", "n0"),
                                  ("action", "Action", "txt"), ("motif", "Motif", "txt"),
                                  ("alertes", "Alertes", "txt")], "Ticker"))
        port = res["portefeuille"]
        p.append(f"<h2>Portefeuille proposé ({len(port)} lignes à poids égal)</h2>")
        p.append(f"<p class='note'>Positions conservées, complétées par les meilleurs signaux ACHETER, "
                 f"au plus {int(cfg['max_par_secteur'] * 100)} % des lignes par secteur.</p>")
        cols = [("action", "Action", "txt"), ("nom", "Société", "txt"), ("secteur_fr", "Secteur", "txt"),
                ("signal", "Signal", "sig"), ("score", "Score", "n0"), ("potentiel_pct", "Potentiel", "n0"),
                ("robustesse_pct", "Robustesse", "n0"), ("mom", "Momentum 12-1", "pct"),
                ("vol", "Volatilité 1 an", "pct"), ("fscore_9", "F-score", "n1"), ("alertes", "Alertes", "txt")]
        p.append(_table_html(port, cols, "Ticker"))
        img = graphique_secteurs(port)
        if img:
            p.append(f'<figure class="petit"><img alt="Secteurs" src="data:image/png;base64,{img}">'
                     "<figcaption>Répartition sectorielle du portefeuille proposé.</figcaption></figure>")
        df = res["classement"]
        cols_c = [("nom", "Société", "txt"), ("secteur_fr", "Secteur", "txt"), ("signal", "Signal", "sig"),
                  ("score", "Score", "n0"), ("potentiel_pct", "Potentiel", "n0"),
                  ("robustesse_pct", "Robustesse", "n0"), ("mom", "Momentum", "pct"),
                  ("revision", "Révision BPA", "pct"), ("rend_fcf", "Rend. FCF", "pct"),
                  ("fscore_9", "F-score", "n1"), ("vol", "Volatilité", "pct"), ("motif", "Motif", "txt")]
        p.append("<h2>Classement : les 60 premiers</h2>")
        p.append(_table_html(df.head(60), cols_c, "Ticker"))
        vendre = df[df["signal"] == "VENDRE"].sort_values("score").head(40)
        p.append("<h2>À éviter : signaux VENDRE (40 plus faibles)</h2>")
        p.append(_table_html(vendre, cols_c, "Ticker"))
        p.append("<p class='note'>Le classement complet (tous les titres, tous les indicateurs) est dans le "
                 "fichier Excel.</p>")
    if bt:
        p.append(section_backtest_html(bt))
    p.append("<h2>Méthode</h2><ul>"
             "<li><b>Filtres de sécurité</b> (éliminatoires) : F-score de Piotroski ≥ 6/9, couverture des intérêts "
             "≥ 3x (1,5x services aux collectivités et immobilier), pas parmi les 10 % les plus endettés de son "
             "secteur, free cash-flow positif sur 3 ans, accruals hors des 10 % les pires du secteur lorsqu'ils "
             "sont positifs (bénéfice supérieur au cash-flow d'exploitation), hors des 10 % les plus volatils de "
             "l'indice. Le M-score de Beneish est signalé en alerte.</li>"
             "<li><b>Potentiel</b> : momentum 12 mois hors dernier mois, dynamique des résultats (variation du "
             "bénéfice trimestriel sur un an, révisions des analystes sur 90 jours), valorisation (rendement "
             "FCF, EBIT / valeur d'entreprise, rendement des bénéfices).</li>"
             "<li><b>Robustesse</b> : qualité (rentabilité brute, ROA, flux de trésorerie, stabilité, "
             "endettement) et faible risque (volatilité, bêta, perte maximale sur un an).</li>"
             "<li>Chaque indicateur est converti en percentile au sein du secteur ; le score final est un "
             "percentile de 0 à 100 sur l'ensemble de l'indice.</li>"
             "<li><b>Signaux</b> : ACHETER = score ≥ 90 et robustesse ≥ 50 ; GARDER = score ≥ 60 ; "
             "SURVEILLER = 50 à 60 ; VENDRE = score &lt; 50 ou échec d'un filtre.</li></ul>")
    p.append("<h2>Limites</h2><ul>"
             "<li>Données gratuites : Yahoo Finance (non officiel) et SEC (chiffres XBRL bruts, parfois "
             "incomplets selon les étiquettes utilisées par chaque société).</li>"
             "<li>Aucun score ne garantit qu'une action ne baissera pas : même l'indice recule en moyenne "
             "d'environ 14 % à un moment de l'année.</li>"
             "<li>Chaque vente peut déclencher de l'impôt sur la plus-value : la zone tampon GARDER / SURVEILLER "
             "sert à limiter les allers-retours.</li></ul>")
    return (f"<!doctype html><html lang='fr'><head><meta charset='utf-8'>"
            f"<meta name='viewport' content='width=device-width, initial-scale=1'>"
            f"<title>Screener S&amp;P 500</title><style>{CSS}</style></head><body><main>"
            + "\n".join(p) + "</main></body></html>")


def exporter_excel(chemin: str, res: Optional[dict], bt: Optional[dict], cfg: Config) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    remplissages = {"ACHETER": "E3F4E3", "GARDER": "E3EEFB", "SURVEILLER": "FDF1D8", "VENDRE": "FBE4E4",
                    "DONNÉES INSUFFISANTES": "EEEDE9", "HORS UNIVERS": "EEEDE9"}

    def mettre_en_forme(ws, df: pd.DataFrame, formats: Dict[str, str], avec_index: bool = True):
        entetes = ([df.index.name or ""] if avec_index else []) + list(df.columns)
        for j, titre in enumerate(entetes, start=1):
            cell = ws.cell(row=1, column=j)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="F0EFEC")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            largeur = 12
            if titre in ("Société", "Motif", "Alertes", "Achats", "Ventes", "Portefeuille", "Action", "Secteur"):
                largeur = {"Société": 28, "Secteur": 24, "Action": 22}.get(titre, 60)
            ws.column_dimensions[get_column_letter(j)].width = largeur
            fmt = formats.get(titre)
            if fmt:
                for i in range(2, ws.max_row + 1):
                    ws.cell(row=i, column=j).number_format = fmt
            if titre == "Signal":
                for i in range(2, ws.max_row + 1):
                    c = ws.cell(row=i, column=j)
                    if c.value in remplissages:
                        c.fill = PatternFill("solid", fgColor=remplissages[c.value])
                        c.font = Font(bold=True)
        ws.freeze_panes = "B2"

    formats = {t: FORMATS_XL[k] for _, t, k in COLONNES_AFFICHAGE if k in FORMATS_XL}
    formats.update({"Poids": "0.0%"})
    with pd.ExcelWriter(chemin, engine="openpyxl") as xw:
        if res:
            t = tableau_affichage(res["classement"])
            t.to_excel(xw, sheet_name="Classement")
            mettre_en_forme(xw.sheets["Classement"], t, formats)
            port = res["portefeuille"].copy()
            tp = tableau_affichage(port)
            tp.insert(0, "Action", port["action"].values)
            tp.insert(1, "Poids", port["poids"].values)
            tp.to_excel(xw, sheet_name="Portefeuille proposé")
            mettre_en_forme(xw.sheets["Portefeuille proposé"], tp, formats)
            if res.get("mes_positions") is not None:
                mp = res["mes_positions"].rename(columns={"ticker": "Ticker", "nom": "Société", "signal": "Signal",
                                                          "score": "Score", "action": "Action", "motif": "Motif",
                                                          "alertes": "Alertes"}).set_index("Ticker")
                mp.to_excel(xw, sheet_name="Mes positions")
                mettre_en_forme(xw.sheets["Mes positions"], mp, {"Score": "0"})
        if bt:
            for nom, tab in bt["metriques"].items():
                feuille = ("BT " + nom.split(" (")[0])[:31]
                tab.index.name = "Indicateur"
                tab.to_excel(xw, sheet_name=feuille)
                mettre_en_forme(xw.sheets[feuille], tab, {})
                ws = xw.sheets[feuille]
                for i, m in enumerate(tab.index, start=2):
                    f = "0.00" if ("Ratio" in m or "Bêta" in m) else "0.0%"
                    for j in range(2, len(tab.columns) + 2):
                        ws.cell(row=i, column=j).number_format = f
            an = bt["annuel"].copy()
            an.index.name = "Année"
            an.to_excel(xw, sheet_name="BT Annuel")
            mettre_en_forme(xw.sheets["BT Annuel"], an, {c: "0.0%" for c in an.columns})
            a = bt["analyse"]
            for nom, tab, fmt in (("BT Facteurs", a.get("ic"), "0.00"), ("BT Quintiles", a.get("quintiles"), "0.0%"),
                                  ("BT Signaux", a.get("signaux"), "0.0%")):
                if tab is not None and not tab.empty:
                    tab.to_excel(xw, sheet_name=nom)
                    mettre_en_forme(xw.sheets[nom], tab, {c: fmt for c in tab.columns})
            j = bt["journal"].rename(columns={"date": "Date", "rebalancement": "Rebalancement",
                                              "positions": "Lignes", "titres": "Titres analysés",
                                              "acheter": "Signaux ACHETER", "insuffisantes": "Données insuffisantes",
                                              "achats": "Achats", "ventes": "Ventes",
                                              "rotation": "Rotation", "frais": "Frais", "liquidites": "Liquidités",
                                              "portefeuille": "Portefeuille"}).set_index("Date")
            j.to_excel(xw, sheet_name="BT Journal")
            mettre_en_forme(xw.sheets["BT Journal"], j, {"Rotation": "0.0%", "Frais": "0.00%", "Liquidités": "0.0%"})
            bt["navs"].rename_axis("Date").to_excel(xw, sheet_name="BT Valeur")
        p = pd.DataFrame({"Valeur": {k: str(v) for k, v in asdict(cfg).items() if k != "email"}})
        p.index.name = "Paramètre"
        p.to_excel(xw, sheet_name="Paramètres")
        xw.sheets["Paramètres"].column_dimensions["A"].width = 28
        xw.sheets["Paramètres"].column_dimensions["B"].width = 40


def metriques_lisibles(tab: pd.DataFrame) -> pd.DataFrame:
    """Version texte (pourcentages, ratios) d'un tableau de métriques, pour l'affichage."""
    out = pd.DataFrame(index=tab.index, columns=tab.columns, dtype=object)
    for m in tab.index:
        k = "n2" if ("Ratio" in m or "Bêta" in m) else "pct"
        out.loc[m] = [_fmt(v, k) for v in tab.loc[m].values]
    return out


def generer_rapports(cfg: Config, res: Optional[dict] = None, bt: Optional[dict] = None) -> Dict[str, str]:
    jour = (res["date"] if res else (bt["fin"] if bt else aujourdhui()))
    suffixe = "" if cfg.mode == "complet" else f"_{cfg.mode}"
    base = cfg.chemin("resultats", f"screener_{pd.Timestamp(jour):%Y-%m-%d}{suffixe}")
    chemins = {"excel": base + ".xlsx", "html": base + ".html"}
    exporter_excel(chemins["excel"], res, bt, cfg)
    with open(chemins["html"], "w", encoding="utf-8") as fh:
        fh.write(rapport_html(res, bt))
    info(f"Rapports enregistrés : {chemins['excel']} et {chemins['html']}")
    return chemins


# ===========================================================================
# 13. Ligne de commande
# ===========================================================================

def lire_positions(arg: str) -> List[str]:
    if not arg:
        return []
    if os.path.exists(arg):
        df = pd.read_csv(arg)
        col = next((c for c in df.columns if str(c).strip().lower() in ("ticker", "symbole", "symbol")), df.columns[0])
        return [ticker_yahoo(t) for t in df[col].dropna().astype(str)]
    return [ticker_yahoo(t) for t in re.split(r"[,\s;]+", arg) if t.strip()]


def main(argv: Optional[Sequence[str]] = None) -> int:
    ap = argparse.ArgumentParser(description="Screener S&P 500 gratuit (Yahoo Finance + SEC EDGAR).")
    ap.add_argument("--email", default=os.environ.get("SEC_EMAIL", ""),
                    help="ton adresse e-mail (exigée par la SEC pour identifier les requêtes)")
    ap.add_argument("--dossier", default="screener_donnees", help="dossier du cache et des résultats")
    ap.add_argument("--positions", default="", help="tes positions : 'AAPL,MSFT' ou un CSV avec une colonne ticker")
    ap.add_argument("--backtest", action="store_true", help="lancer aussi le backtest")
    ap.add_argument("--backtest-seul", action="store_true", help="lancer uniquement le backtest")
    ap.add_argument("--debut", default=None, help="début du backtest (AAAA-MM-JJ)")
    ap.add_argument("--mode", choices=["complet", "prix"], default="complet")
    ap.add_argument("--sans-revisions", action="store_true", help="ne pas interroger les révisions d'analystes")
    ap.add_argument("--reglages", default=None, help="fichier JSON de réglages (celui de l'application web)")
    a = ap.parse_args(argv)
    base = {}
    if a.reglages:
        with open(a.reglages, encoding="utf-8") as fh:
            base = json.load(fh)
    cfg = Config.depuis_dict(base, email=a.email or "ton.email@exemple.com", dossier=a.dossier)
    if a.mode != "complet" or "mode" not in base:
        cfg.mode = a.mode
    if a.sans_revisions:
        cfg.revisions_analystes = False
    if a.debut:
        cfg.debut_backtest = a.debut
    if cfg.mode == "complet" and not cfg.email_valide():
        print("Erreur : indique ton adresse e-mail avec --email (la SEC l'exige), ou utilise --mode prix.")
        return 2
    res = None if a.backtest_seul else lancer_screener(cfg, lire_positions(a.positions))
    bt = lancer_backtest(cfg) if (a.backtest or a.backtest_seul) else None
    generer_rapports(cfg, res, bt)
    if res is not None:
        print("\nPortefeuille proposé :")
        print(res["portefeuille"][["action", "nom", "secteur_fr", "signal", "score"]].to_string())
    if bt is not None:
        for nom, tab in bt["metriques"].items():
            print(f"\n{nom}\n" + tab.round(3).to_string())
    return 0


if __name__ == "__main__":
    sys.exit(main())
